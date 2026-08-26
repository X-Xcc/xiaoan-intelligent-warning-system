from pathlib import Path

from pptx import Presentation


PPTX = Path(r"D:\xx\Desktop\烟火PPT.pptx")
OUT = Path(r"D:\CICSIC\ppt_analysis\yanhuo_slide_text.md")


def shape_text(shape):
    if not hasattr(shape, "text"):
        return ""
    return "\n".join(line.strip() for line in shape.text.splitlines() if line.strip())


def main():
    prs = Presentation(PPTX)
    lines = [f"# {PPTX.name}", "", f"Slides: {len(prs.slides)}", ""]
    for idx, slide in enumerate(prs.slides, 1):
        lines.append(f"## Slide {idx}")
        items = []
        for shape in slide.shapes:
            text = shape_text(shape)
            if not text:
                continue
            items.append((shape.top, shape.left, text))
        if not items:
            lines.append("(no extractable text)")
        else:
            for _, _, text in sorted(items):
                lines.append(text)
                lines.append("")
        lines.append("")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
