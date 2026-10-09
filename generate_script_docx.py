from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r'D:\CICSIC\烟火哨兵七分钟路演逐字稿.docx')
SPOKEN = []

def font(run, east='仿宋', latin='Times New Roman', size=15, bold=False):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn('w:eastAsia'), east)
    rpr.rFonts.set(qn('w:ascii'), latin)
    rpr.rFonts.set(qn('w:hAnsi'), latin)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def spacing(p, before=0, after=4, line=1.45):
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = line

def page_field(p):
    r = p.add_run()
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
    b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = 'PAGE'
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'end')
    r._r.extend([a, b, c]); font(r, size=10)

def heading(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; spacing(p, before=8, after=4, line=1.2)
    font(p.add_run(text), east='黑体', size=16, bold=True)

def flip(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.74); p.paragraph_format.keep_with_next = True; spacing(p, after=3, line=1.2)
    font(p.add_run(text), east='黑体', size=12, bold=True)

def script(doc, text):
    SPOKEN.append(text)
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0.74); p.paragraph_format.keep_together = True; spacing(p)
    font(p.add_run(text)); return p

doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.2); sec.left_margin = Cm(2.6); sec.right_margin = Cm(2.4)
normal = doc.styles['Normal']; normal.font.name = 'Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋'); normal.font.size = Pt(15)
normal.paragraph_format.widow_control = True
header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT; font(header.add_run('烟火哨兵项目路演逐字稿'), size=9)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(footer.add_run('第 '), size=10); page_field(footer); font(footer.add_run(' 页'), size=10)

p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, before=18, after=8, line=1.2); font(p.add_run('烟火哨兵项目七分钟路演逐字稿'), east='黑体', size=22, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=18, line=1.2); font(p.add_run('公共安全智能预警与秒级响应系统'), size=15)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=16, line=1.3); font(p.add_run('适用场景：项目路演、创新大赛答辩    建议时长：约7分钟'), size=12)
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.74); p.paragraph_format.right_indent = Cm(0.74); spacing(p, after=12, line=1.35)
font(p.add_run('使用提示：'), east='黑体', size=12, bold=True); font(p.add_run('仅朗读正文，标题和翻页提示不念。按PPT顺序展示，同组页面随讲述推进。时长按每分钟约300字并配合自然停顿估算，以现场排练为准。'), size=12)

heading(doc, '一 项目概述')
flip(doc, '[翻页：第1至2页]')
script(doc, '各位评委好，我们是江西司法警官职业学院“烟火哨兵”项目团队。今天汇报的项目是公共安全智能预警与秒级响应系统。我们聚焦夜市治安场景，将智能预警、动态识别与快速处置相结合，推动夜市治理从人防向智防升级，保障群众消费安全与商户营商安全。项目已获得省职业院校技能大赛一等奖。')

heading(doc, '二 建设背景与需求调研')
flip(doc, '[翻页：第3至6页]')
script(doc, '项目源于夜市一线的治理需求。夜市人员密集，酒后滋事、摊位纠纷和财物盗窃等事件影响经营秩序。真实案例显示，细小的冲突苗头可能迅速升级，因此，风险发现和干预需要前移。社会治安整体防控与“互联网加公共安全”建设，为项目提供了政策方向。目前，南昌规范化夜市63处，江西200余处，全国5000余处，数字化治理需求持续释放。')
flip(doc, '[翻页：第7至10页]')
script(doc, '围绕这一需求，我们结合文献研究与实地调研，覆盖基层公安、商户、消费者和夜市管理中心四类主体，开展8场走访，发放问卷1500份，回收有效问卷959份。64.52%的商户反馈险情发生后人员抵达存在延迟，67.74%的商户希望优先处置醉酒滋事骚扰事件，47.42%的消费者最担忧财物被盗。调研据此归纳出预警弱、识别难、响应慢三项治理痛点，也明确了项目的技术方向。')

heading(doc, '三 系统方案与核心技术')
flip(doc, '[翻页：第11至13页]')
script(doc, '针对这三项痛点，烟火哨兵形成前置管控、行为甄别、快速处置的全流程治安闭环，连接管理控制台、监控大屏与小程序。自2025年10月启动实地调研以来，团队逐步完成模块研发、试点验证和系统集成，2026年8月推进整套系统部署与迭代优化。具体通过三项核心技术实现。')
flip(doc, '[翻页：第14页]')
script(doc, '第一，跨夜市分级研判预警。商户和群众上报滋扰线索，经公安核验入库后，系统关联历史数据进行风险分级：3次以下留存事件触发风险提示，3次及以上启动重点管控，实现跨夜市风险信息联动，将处置关口前移。')
flip(doc, '[翻页：第15至16页]')
script(doc, '在前置研判基础上，第二项多模态风险识别技术进一步关注现场行为。机器狗采集移动视角视频，由YOLOv8初筛，再由千问视觉大模型复核，识别推搡、纠缠、拉扯等行为，并从汇聚增速、群体密度和滞留时长分析聚集风险。双层识别将系统误报率控制在4%至6%，相较传统方式下降75%至80%，为冲突萌芽阶段的干预提供依据。')
flip(doc, '[翻页：第17至18页]')
script(doc, '发现风险后，第三项线上一键报警与无人机联动响应技术接续处置。商户通过小程序一键报警，警情与位置直达属地派出所和巡逻警力，同时调度无人机喊话、图传和留证。小程序经过6版迭代、300余小时打磨，平台上报平均定位偏差由口述方式的80米降至5米。无人机60秒抵达现场上空，可节省约3分钟前置处置时间。')

heading(doc, '四 研发成果与应用验证')
flip(doc, '[翻页：第19至22页]')
script(doc, '围绕上述技术，项目取得4项已授权软件著作权，2项发明专利申请处于受理阶段，并获得警务实战专家认可。第三方机构开展63项检测，结果显示核心能力达标、响应达到毫秒级、安全防护可靠。科技查新检索19个数据库，结论为国内未见相同文献报道，为项目创新性提供依据。')
flip(doc, '[翻页：第23页]')
script(doc, '技术成果已进入基层试点。桃花派出所开展30天分级研判预警应用，场所纠纷萌芽期干预率提升80%，夜市风险前置识别能力提升75%，滋扰类纠纷警情压降12%。')
flip(doc, '[翻页：第24至25页]')
script(doc, '荆公路派出所开展98天一键报警试点，抵达现场时间缩短3至5分钟，人力资源占用降低50%。九洲、绳金塔派出所开展多模态风险识别试点，由机器狗辅助识别危险动作，巡逻视野盲区降低12%。四家派出所的试点应用，分别验证预警、报警联动和现场识别功能。')

heading(doc, '五 商业模式与落地计划')
flip(doc, '[翻页：第26至28页]')
script(doc, '以试点应用为基础，项目面向基层公安机关、政府和安保公司，采用轻量试点、接入存量设备的落地路径，联合硬件与云服务供应商，通过样板推广、科研合作和线上线下宣传拓展市场。')
flip(doc, '[翻页：第29页]')
script(doc, '收入由软件服务、硬件销售、平台增值和技术培训构成。软件服务定价每年30万元，硬件每套8万元，软件服务预计占收入的67%，形成持续运营收入。')
flip(doc, '[翻页：第30至31页]')
script(doc, '财务预测显示，2027年至2030年营业收入分别为170万元、320万元、480万元和700万元，累计净利润约500万元，预计2028年达到盈亏平衡。团队计划融资100万元，出让15%股权，资金用于研发、市场推广、团队建设和运营储备，并计划于2027年1月成立公司。')

heading(doc, '六 团队资源与实践成长')
flip(doc, '[翻页：第32至38页]')
script(doc, '项目实施依托7人跨专业团队，成员来自司法信息技术、治安管理、无人机应用技术和大数据会计专业，协同承担开发、训练、调研、硬件与运营工作。负责人具有基层派出所实习和系统研发经历，指导教师与行业专家持续提供技术、警务和商业指导。')
flip(doc, '[翻页：第39至43页]')
script(doc, '在共同研发中，团队将任务分解为七个模块，通过协同调试和多轮优化，把专业知识转化为实际应用能力。学校通过校局合作提供真实场景与实训资源，项目经验进一步转化为教学案例，推动人才培养与基层警务需求衔接。')

heading(doc, '七 社会价值与发展规划')
flip(doc, '[翻页：第44至47页]')
script(doc, '项目将以智能安防促进安保行业数字化转型，带动硬件、算法与运维发展，已获得中国新闻资讯网、江西新闻网等媒体报道。前期计划带动8个岗位，未来三年计划间接带动约50人就业。通过风险早发现、快处置，提升群众安全感，为夜间经济发展提供保障。')
flip(doc, '[翻页：第48页]')
script(doc, '未来，2026年至2027年以本市试点夯实基础，2028年至2029年推进省域拓展，目标是业务覆盖江西、服务超过500万人、营业收入突破300万元。2030年及以后，在省内规模化应用基础上逐步拓展全国市场。')
flip(doc, '[翻页：第49页]')
script(doc, '以科技护烟火，以闭环守平安。烟火哨兵将持续深化夜市商圈智慧治安实践，为夜间经济安全治理提供可复制的经验。我的汇报完毕，谢谢各位评委。')

doc.save(OUT)
print(OUT)
print('spoken_characters=', sum(len(t) for t in SPOKEN))
print('spoken_han=', sum('\u4e00' <= c <= '\u9fff' for t in SPOKEN for c in t))
