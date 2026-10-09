from zipfile import ZipFile
from docx import Document

path = r"C:\Users\xx\Desktop\烟火哨兵七分钟路演逐字稿.docx"
doc = Document(path)
print("chars", sum(len(p.text) for p in doc.paragraphs))
print("flips", sum(1 for p in doc.paragraphs if p.text.startswith("[翻页")))
print("paragraphs", len(doc.paragraphs))
print("zip_ok", ZipFile(path).testzip() is None)
