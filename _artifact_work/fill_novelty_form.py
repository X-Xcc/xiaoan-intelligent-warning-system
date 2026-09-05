from pathlib import Path
import json
import shutil

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


SRC = Path(r"D:\xx\Desktop\查新委托单(1)(12)(1).docx")
OUT = Path(r"D:\CICSIC\_artifact_work\查新委托单-城脉智诊已填写-定稿.docx")
AUDIT = Path(r"D:\CICSIC\_artifact_work\查新委托单-城脉智诊填写核对.json")


PROJECT_CN = "CS城脉智诊-无监督驱动的地下管网分布式智能监测与边缘预警系统"
PROJECT_EN = (
    "CS Chengmai Zhizhen Unsupervised Distributed Smart Monitoring and Edge "
    "Early Warning System for Underground Pipe Networks"
)

TECH_POINTS = (
    "本项目属于地下管网分布式智能监测与边缘预警技术领域，面向管网泄漏、压力异常、堵塞、爆管、"
    "微漏、暗漏和微破损等异常工况的识别、定位与处置需求，构建由分布式传感节点、边缘网关和平台"
    "联动机制组成的智能监测系统。系统采集声学、压力、流量等多源数据，完成统一解析、融合分析、"
    "异常感知、类型识别、空间定位和分级预警，并将异常点与管网 GIS 图层匹配，实现预警工单自动"
    "派发、流转、处置核验和异常数据溯源。检测材料显示，系统覆盖 10 类常见管网异常工况，在无需"
    "标注样本条件下快速适配；异常事件感知精度为 98.7%，误报率为 1.1%，漏报率为 0.3%，异常类型"
    "识别覆盖率为 98.2%，新增工况 48 小时内完成模型自主优化；端到端预警响应时间为 3.2 秒，异常"
    "点空间定位误差为 0.68 米，边缘端离线可独立运行 96 小时。"
)

NOVELTY_POINTS = (
    "1. 无监督驱动的地下管网多源数据融合监测与异常类型识别方法，支持声学、压力、流量数据统一解析与快速适配。\n"
    "2. 面向地下管网异常的边缘端实时预警与空间定位机制，实现秒级响应、米级定位和多点并发异常识别。\n"
    "3. 具备离线独立运行、自动派单、处置核验、GIS 匹配和长期本地存储的分布式闭环运维机制。"
)

KEYWORDS = (
    "中文检索词：地下管网；分布式监测；无监督学习；异常检测；边缘预警；多源传感数据融合；"
    "声学传感；压力传感；流量传感；泄漏识别；堵塞识别；爆管预警；故障定位；GIS 匹配；离线运行；"
    "闭环运维。\n"
    "英文检索词：unsupervised learning; distributed monitoring; underground pipe network; edge early warning; "
    "multi-source sensor fusion; acoustic sensing; pressure sensing; flow sensing; anomaly detection; "
    "fault localization; GIS matching; leakage; blockage; pipe burst; offline operation."
)

IP_STATUS = (
    "暂无"
)

REFERENCES = (
    "1. 深圳市中安质量检验认证有限公司：《城脉智诊-无监督驱动的地下管网分布式智能监测与边缘预警系统"
    "检验报告》，报告编号 CTS26315A824B23853YG，2026年05月18日。\n"
    "2. GB/T 4208-2017；GB/T 51354-2019；GB/T 2423 系列；GB/T 9254.1-2021。"
)


def clear_cell(cell):
    for p in cell.paragraphs:
        p._element.getparent().remove(p._element)


def add_paragraph(cell, text="", bold=False, size=10.5, align=None, first_line=False):
    p = cell.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if first_line:
        p.paragraph_format.first_line_indent = Pt(21)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    return p


def set_cell_text(cell, text, bold=False, size=10.5, align=None):
    clear_cell(cell)
    parts = str(text).split("\n")
    for i, part in enumerate(parts):
        add_paragraph(cell, part, bold=bold, size=size, align=align)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_section_cell(cell, title, body, body_size=9.8):
    clear_cell(cell)
    add_paragraph(cell, title, bold=True, size=10.5)
    for line in body.split("\n"):
        add_paragraph(cell, line, size=body_size, first_line=not line[:2].isdigit())
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def normalize_table_rows(table):
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)

    doc = Document(OUT)
    table = doc.tables[0]
    normalize_table_rows(table)

    set_cell_text(table.rows[0].cells[2], PROJECT_CN, size=11)
    set_cell_text(table.rows[1].cells[2], PROJECT_EN, size=9.2)
    set_cell_text(table.rows[2].cells[2], "沈阳大学", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.rows[3].cells[2], "沈阳市大东区望花南街21号", size=11)
    set_cell_text(table.rows[4].cells[2], "清池", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.rows[4].cells[4], "13380185668", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    purpose = (
        "二、查新目的及范围\n"
        "1.查新目的：(必填)\n"
        "√产品查新\n"
        "2.查新范围：(必填)\n"
        "√国内外查新"
    )
    set_cell_text(table.rows[6].cells[0], purpose, size=10.5)
    set_section_cell(table.rows[7].cells[0], "三、查新项目的科学技术要点 (必填)", TECH_POINTS, body_size=8.8)
    set_section_cell(table.rows[8].cells[0], "四、查新点(必填)", NOVELTY_POINTS, body_size=9.3)
    set_section_cell(table.rows[9].cells[0], "五、参考检索词及其解释", KEYWORDS, body_size=8.8)
    set_section_cell(table.rows[10].cells[0], "六、知识产权及已发表论文情况", IP_STATUS, body_size=9.3)
    set_section_cell(table.rows[11].cells[0], "七、参考文献", REFERENCES, body_size=9.0)

    for section in doc.sections:
        section.top_margin = section.top_margin
        section.bottom_margin = section.bottom_margin

    cp = doc.core_properties
    cp.title = "科技查新委托单"
    cp.subject = PROJECT_CN
    cp.keywords = "科技查新;地下管网;无监督学习;边缘预警"

    doc.save(OUT)

    check_doc = Document(OUT)
    check_table = check_doc.tables[0]
    values = {
        "中文项目名称": check_table.rows[0].cells[2].text,
        "英文项目名称": check_table.rows[1].cells[2].text,
        "委托单位": check_table.rows[2].cells[2].text,
        "收件地址": check_table.rows[3].cells[2].text,
        "收件人": check_table.rows[4].cells[2].text,
        "手机": check_table.rows[4].cells[4].text,
        "查新目的及范围": check_table.rows[6].cells[0].text,
        "科学技术要点": check_table.rows[7].cells[0].text,
        "查新点": check_table.rows[8].cells[0].text,
        "参考检索词": check_table.rows[9].cells[0].text,
        "知识产权及论文": check_table.rows[10].cells[0].text,
        "参考文献": check_table.rows[11].cells[0].text,
    }
    placeholders = ["国内查新必填", "国内外查新必填", "必填"]
    unresolved = {}
    for name, value in values.items():
        hits = [p for p in placeholders if p in value]
        if hits:
            unresolved[name] = hits
    AUDIT.write_text(
        json.dumps(
            {
                "output": str(OUT),
                "filled_fields": values,
                "unresolved_required_placeholders": unresolved,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(str(OUT))
    print(str(AUDIT))
    print(json.dumps({"unresolved_required_placeholders": unresolved}, ensure_ascii=False))


if __name__ == "__main__":
    main()
