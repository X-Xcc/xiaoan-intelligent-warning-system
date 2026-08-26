from pathlib import Path
import sys
from zipfile import ZipFile
import re

from pptx import Presentation


def main():
    desktop = Path(r"D:\xx\Desktop")
    if len(sys.argv) > 1:
        p = Path(sys.argv[1])
    else:
        p = desktop / "新建 PPTX 演示文稿.pptx"
        if not p.exists():
            files = sorted(
                [p for p in desktop.glob("*.pptx") if not p.name.startswith("~$")],
                key=lambda p: p.stat().st_mtime,
                reverse=True,
            )
            if not files:
                raise SystemExit("no pptx files found")
            p = files[0]
    print(f"FILE {p}")
    with ZipFile(p) as z:
        slides = sorted(
            [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")],
            key=lambda s: int(Path(s).stem.replace("slide", "")),
        )
        print(f"SLIDES {len(slides)}")
        for s in slides:
            idx = int(Path(s).stem.replace("slide", ""))
            xml = z.read(s).decode("utf-8", errors="ignore")
            texts = re.findall(r"<a:t>(.*?)</a:t>", xml)
            print(f"---SLIDE {idx}---")
            print(" | ".join(texts).encode("unicode_escape").decode("ascii"))

    print("=== SHAPES ===")
    prs = Presentation(str(p))
    def walk_shapes(shapes, prefix=""):
        for sidx, shape in enumerate(shapes, start=1):
            label = f"{prefix}[{sidx}] {shape.name} {shape.shape_type}"
            if shape.shape_type == 6:  # GROUP
                print(label)
                walk_shapes(shape.shapes, prefix + "  ")
                continue
            txt = ""
            if getattr(shape, "has_text_frame", False):
                txt = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text)
            elif getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    rows.append(" || ".join(cell.text for cell in row.cells))
                txt = "\n".join(rows)
            if txt:
                print(label)
                print(txt.encode("unicode_escape").decode("ascii"))

    for idx, slide in enumerate(prs.slides, start=1):
        print(f"---SLIDE {idx} SHAPES---")
        walk_shapes(slide.shapes)


if __name__ == "__main__":
    main()
