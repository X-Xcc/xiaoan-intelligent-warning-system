from pathlib import Path
from zipfile import ZipFile
from collections import Counter
from docx import Document
from pypdf import PdfReader
from pdf2image import convert_from_path
from PIL import Image, ImageDraw
import json
import re

root = Path(__file__).parent
source = root / 'source_snapshot.docx'
output = Path(r'D:/xx/Desktop/烟火哨兵——公共安全智能预警与秒级响应系统_去AI味修订版.docx')
original, revised = Document(source), Document(output)
changes = json.loads(root.joinpath('changes.json').read_text(encoding='utf-8'))
expected = [paragraph.text for paragraph in original.paragraphs]
for change in changes:
    expected[change['paragraph']] = change['after']
assert expected == [paragraph.text for paragraph in revised.paragraphs]
assert Counter(re.findall(r'\d+(?:[.,]\d+)*', '\n'.join(expected))) == Counter(re.findall(r'\d+(?:[.,]\d+)*', '\n'.join(paragraph.text for paragraph in original.paragraphs)))
assert [[cell.text for row in table.rows for cell in row.cells] for table in original.tables] == [[cell.text for row in table.rows for cell in row.cells] for table in revised.tables]
with ZipFile(source) as before, ZipFile(output) as after:
    before_media = Counter(before.read(name) for name in before.namelist() if name.startswith('word/media/') and not name.endswith('/'))
    after_media = Counter(after.read(name) for name in after.namelist() if name.startswith('word/media/') and not name.endswith('/'))
    assert before_media == after_media
for old_section, new_section in zip(original.sections, revised.sections):
    old_backgrounds = old_section.header._element.xpath('.//wp:anchor')
    new_backgrounds = new_section.header._element.xpath('.//wp:anchor')
    assert len(old_backgrounds) == len(new_backgrounds)
    for old_anchor, new_anchor in zip(old_backgrounds, new_backgrounds):
        assert old_anchor.get('behindDoc') == new_anchor.get('behindDoc')
        assert [(extent.get('cx'), extent.get('cy')) for extent in old_anchor.xpath('./wp:extent')] == [(extent.get('cx'), extent.get('cy')) for extent in new_anchor.xpath('./wp:extent')]
pdf = root / 'qa.pdf'
reader = PdfReader(pdf)
pages = [{'page': index + 1, 'chars': len(page.extract_text().strip())} for index, page in enumerate(reader.pages)]
images = convert_from_path(pdf, size=1400, poppler_path=r'C:/Users/xx/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin', thread_count=4)
for index, image in enumerate(images):
    image.save(root / f'page-{index + 1:02d}.png')
for start in range(0, len(images), 16):
    sheet = Image.new('RGB', (1920, 2820), '#cccccc')
    draw = ImageDraw.Draw(sheet)
    for index, image in enumerate(images[start:start + 16]):
        preview = image.copy()
        preview.thumbnail((470, 665))
        left, top = (index % 4) * 480, (index // 4) * 705
        sheet.paste(preview, (left, top + 25))
        draw.text((left + 8, top + 5), f'PAGE {start + index + 1}', fill='black')
    sheet.save(root / f'contact-{start // 16 + 1}.jpg')
audit = {'skill': 'humanizer-zh', 'paragraphs_rewritten': len(changes), 'paragraph_count': len(revised.paragraphs), 'page_count': len(reader.pages), 'characters_before': sum(len(paragraph.text) for paragraph in original.paragraphs), 'characters_after': sum(len(paragraph.text) for paragraph in revised.paragraphs), 'all_numeric_tokens_preserved': True, 'tables_unchanged': True, 'media_unchanged': True, 'page_backgrounds_preserved': True, 'pages': pages}
root.joinpath('audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({key: value for key, value in audit.items() if key != 'pages'}, ensure_ascii=False))
print('Sparse pages:', [page for page in pages if page['chars'] < 100])
