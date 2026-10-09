from pathlib import Path
from pdf2image import convert_from_path
from PIL import Image, ImageDraw
from pypdf import PdfReader
from zipfile import ZipFile
from collections import Counter
from docx import Document

root = Path(__file__).parent
source = Path(r'D:/xx/Desktop/烟火哨兵——公共安全智能预警与秒级相应系统.docx')
output = source.with_name('烟火哨兵——公共安全智能预警与秒级响应系统_沿用原封面纸张美化版.docx')
pdf = output.with_suffix('.pdf')
reader = PdfReader(pdf)
print('Pages:', len(reader.pages))
print('Sparse pages:', [(index + 1, len(page.extract_text())) for index, page in enumerate(reader.pages) if len(page.extract_text().strip()) < 100])
with ZipFile(source) as original, ZipFile(output) as polished:
    original_media = Counter(original.read(name) for name in original.namelist() if name.startswith('word/media/') and not name.endswith('/'))
    output_media = Counter(polished.read(name) for name in polished.namelist() if name.startswith('word/media/') and not name.endswith('/'))
    assert not original_media - output_media
original, polished = Document(source), Document(output)
original_paragraphs = [paragraph.text for paragraph in original.paragraphs if paragraph.text.strip()]
output_paragraphs = [paragraph.text for paragraph in polished.paragraphs if paragraph.text.strip() and not paragraph.style.name.startswith('TOC')]
assert original_paragraphs == output_paragraphs
assert [[cell.text for row in table.rows for cell in row.cells] for table in original.tables] == [[cell.text for row in table.rows for cell in row.cells] for table in polished.tables]
print('All original body text, table data, and media preserved')
images = convert_from_path(pdf, size=650, poppler_path=r'C:/Users/xx/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin', thread_count=4)
for start in range(0, len(images), 16):
    sheet = Image.new('RGB', (1920, 2820), '#cccccc')
    draw = ImageDraw.Draw(sheet)
    for index, image in enumerate(images[start:start + 16]):
        image.thumbnail((470, 665))
        left, top = (index % 4) * 480, (index // 4) * 705
        sheet.paste(image, (left, top + 25))
        draw.text((left + 8, top + 5), f'PAGE {start + index + 1}', fill='black')
    sheet.save(root / f'final-contact-{start // 16 + 1}.jpg')
images[0].save(root / 'final-cover.png')
