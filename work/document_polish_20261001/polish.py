from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

source = Path(r'D:/xx/Desktop/烟火哨兵——公共安全智能预警与秒级相应系统.docx')
output = source.with_name('烟火哨兵——公共安全智能预警与秒级响应系统_封面页眉页脚美化版.docx')
document = Document(source)
original_text = [paragraph.text for paragraph in document.paragraphs]
original_tables = [[cell.text for row in table.rows for cell in row.cells] for table in document.tables]

def font_set(run, size, name='宋体', bold=None, color='000000'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)
    if bold is not None:
        run.bold = bold

def property_set(parent, tag, attributes):
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    for key, value in attributes.items():
        element.set(qn(key), str(value))
    return element

for section in document.sections:
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin, section.right_margin = Cm(2.8), Cm(2.2)
    section.header_distance = section.footer_distance = Cm(1.3)
for index, paragraph in enumerate(document.paragraphs):
    formatting = paragraph.paragraph_format
    formatting.widow_control = True
    if index < 8:
        continue
    if paragraph.style.name.startswith('Heading'):
        level = int(paragraph.style.name.split()[-1])
        formatting.first_line_indent = Pt(0)
        formatting.keep_with_next = formatting.keep_together = True
        formatting.line_spacing = 1.25
        formatting.space_before = Pt(0 if level == 1 else 12 if level == 2 else 8)
        formatting.space_after = Pt(16 if level == 1 else 6 if level == 2 else 4)
        formatting.alignment = Align.CENTER if level == 1 else Align.LEFT
        if level == 1:
            formatting.page_break_before = True
        for run in paragraph.runs:
            font_set(run, {1: 16, 2: 14, 3: 12, 4: 12}[level], '黑体', True)
    elif paragraph._p.xpath('.//w:drawing'):
        formatting.alignment = Align.CENTER
        formatting.first_line_indent = Pt(0)
        formatting.keep_with_next = formatting.keep_together = True
        formatting.line_spacing = 1.0
        formatting.space_before, formatting.space_after = Pt(8), Pt(3)
    elif re.match(r'^[图表]\s*\d+[.．-]\d+', paragraph.text.strip()):
        formatting.alignment = Align.CENTER
        formatting.first_line_indent = Pt(0)
        formatting.line_spacing = 1.2
        formatting.space_before, formatting.space_after = Pt(3), Pt(8)
        formatting.keep_together = True
        formatting.keep_with_next = paragraph.text.lstrip().startswith('表')
        for run in paragraph.runs:
            font_set(run, 10.5, color='595959')
    elif paragraph.text.strip():
        formatting.alignment = Align.JUSTIFY
        formatting.line_spacing = Pt(20)
        formatting.space_before, formatting.space_after = Pt(0), Pt(0)
        formatting.first_line_indent = Pt(24)
        formatting.keep_with_next = False
        if re.match(r'^\[\d+\]', paragraph.text.strip()):
            formatting.first_line_indent, formatting.left_indent = Pt(-21), Pt(21)
            formatting.line_spacing = 1.35
        for run in paragraph.runs:
            font_set(run, 12)
specs = [(18, '黑体', 12, 10), (16, '宋体', 0, 24), (34, '黑体', 0, 12), (17, '黑体', 0, 12), (10, '宋体', 0, 12), (12, '宋体', 8, 0)]
for index, (size, name, before, after) in enumerate(specs):
    paragraph = document.paragraphs[index]
    paragraph.paragraph_format.alignment = Align.CENTER
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(before), Pt(after)
    if index == 2:
        paragraph.style = document.styles['Title']
    for run in paragraph.runs:
        font_set(run, size, name, index in [0, 2], '24516D' if index == 3 else '595959' if index in [4, 5] else '000000')
for level in range(1, 4):
    name = f'TOC {level}'
    style = document.styles[name] if name in document.styles else document.styles.add_style(name, 1)
    style.font.name, style.font.size = 'Times New Roman', Pt(12 if level < 3 else 11)
    style.font.color.rgb, style.font.bold = RGBColor(0, 0, 0), level == 1
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体' if level == 1 else '宋体')
    style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(6 if level == 1 else 0), Pt(3)
    style.paragraph_format.line_spacing = 1.2
    style.paragraph_format.left_indent, style.paragraph_format.first_line_indent = Pt((level - 1) * 12), Pt(0)
for table_index, table in enumerate(document.tables):
    if table_index == 0:
        for row in table.rows:
            for cell_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(5)
                    paragraph.paragraph_format.line_spacing = 1.2
                    for run in paragraph.runs:
                        font_set(run, 12, '黑体' if cell_index == 0 else '宋体', cell_index == 0, '24516D' if cell_index == 0 else '333333')
                    property_set(cell._tc.get_or_add_tcPr(), 'w:shd', {'w:fill': 'EEF3F7', 'w:val': 'clear'})
        continue
    table.autofit = False
    borders = property_set(table._tbl.tblPr, 'w:tblBorders', {})
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        property_set(borders, f'w:{side}', {'w:val': 'single', 'w:sz': 4, 'w:color': 'C9D1D9'})
    for row_index, row in enumerate(table.rows):
        property_set(row._tr.get_or_add_trPr(), 'w:cantSplit', {})
        if row_index == 0:
            property_set(row._tr.get_or_add_trPr(), 'w:tblHeader', {})
        for cell_index, cell in enumerate(row.cells):
            properties = cell._tc.get_or_add_tcPr()
            property_set(properties, 'w:vAlign', {'w:val': 'center'})
            property_set(properties, 'w:shd', {'w:fill': 'E4EBF2' if row_index == 0 else 'F5F7FA' if row_index % 2 == 0 else 'FFFFFF', 'w:val': 'clear'})
            margins = property_set(properties, 'w:tcMar', {})
            for side, value in [('top', 80), ('bottom', 80), ('left', 90), ('right', 90)]:
                property_set(margins, f'w:{side}', {'w:w': value, 'w:type': 'dxa'})
            for paragraph in cell.paragraphs:
                formatting = paragraph.paragraph_format
                formatting.first_line_indent, formatting.line_spacing = Pt(0), 1.15
                formatting.space_before = formatting.space_after = Pt(1)
                formatting.keep_with_next = row_index == 0
                formatting.alignment = Align.CENTER if row_index == 0 or (len(table.columns) == 5 and cell_index > 0) else Align.LEFT
                for run in paragraph.runs:
                    font_set(run, 10.5, '黑体' if row_index == 0 else '宋体', row_index == 0)
assert original_text == [paragraph.text for paragraph in document.paragraphs]
assert original_tables == [[cell.text for row in table.rows for cell in row.cells] for table in document.tables]
cover_image = deepcopy(document.paragraphs[13]._p)
for extent in cover_image.xpath('.//wp:extent | .//a:xfrm/a:ext'):
    extent.set('cx', str(int(Cm(14.4))))
    extent.set('cy', str(int(Cm(7.1))))
for properties in cover_image.xpath('.//wp:docPr'):
    properties.set('id', '5001')
    properties.set('name', 'Cover Project Illustration')
document.tables[0]._tbl.addprevious(cover_image)
for section in document.sections:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = Align.LEFT
    header.paragraph_format.space_after = Pt(6)
    tabs = property_set(header._p.get_or_add_pPr(), 'w:tabs', {})
    property_set(tabs, 'w:tab', {'w:val': 'right', 'w:pos': 9072})
    font_set(header.add_run('烟火哨兵  SMOKE WATCH'), 10, '黑体', True, '24516D')
    font_set(header.add_run('\t公共安全智能预警与秒级响应系统'), 9, color='697B88')
    borders = property_set(header._p.get_or_add_pPr(), 'w:pBdr', {})
    property_set(borders, 'w:bottom', {'w:val': 'single', 'w:sz': 10, 'w:space': 7, 'w:color': '24516D'})
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = Align.LEFT
    footer.paragraph_format.space_before = Pt(6)
    tabs = property_set(footer._p.get_or_add_pPr(), 'w:tabs', {})
    center_tab = OxmlElement('w:tab')
    center_tab.set(qn('w:val'), 'center')
    center_tab.set(qn('w:pos'), '4536')
    tabs.append(center_tab)
    right_tab = OxmlElement('w:tab')
    right_tab.set(qn('w:val'), 'right')
    right_tab.set(qn('w:pos'), '9072')
    tabs.append(right_tab)
    font_set(footer.add_run('大学生创新创业项目策划书\t'), 8.5, color='697B88')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    font_set(footer.add_run('\t江西司法警官职业学院'), 8.5, color='697B88')
    borders = property_set(footer._p.get_or_add_pPr(), 'w:pBdr', {})
    property_set(borders, 'w:top', {'w:val': 'single', 'w:sz': 4, 'w:space': 7, 'w:color': 'AABBC8'})
main_section = document.sections[1]
header_reference = deepcopy(main_section._sectPr.find(qn('w:headerReference')))
footer_reference = deepcopy(main_section._sectPr.find(qn('w:footerReference')))
for section_properties in document.element.xpath('.//w:sectPr'):
    for tag, replacement in [('w:headerReference', header_reference), ('w:footerReference', footer_reference)]:
        for previous in list(section_properties.findall(qn(tag))):
            section_properties.remove(previous)
        section_properties.insert(0, deepcopy(replacement))
document.save(output)
with ZipFile(source) as original, ZipFile(output) as polished:
    media = [name for name in original.namelist() if name.startswith('word/media/') and not name.endswith('/')]
    assert all(original.read(name) == polished.read(name) for name in media)
print(output)
print(f'Preserved {len(original_text)} paragraphs, {len(original_tables)} tables, {len(media)} media files')
