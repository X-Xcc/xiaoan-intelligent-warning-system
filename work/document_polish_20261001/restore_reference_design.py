from pathlib import Path
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from collections import Counter
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as Relations
from docx.parts.hdrftr import HeaderPart

source = Path(r'D:/xx/Desktop/烟火哨兵——公共安全智能预警与秒级相应系统.docx')
reference_path = Path(r'D:/xx/Desktop/烟火哨兵  材料收集表/烟火哨兵/烟火哨兵——公共安全智能预警与秒级响应系统.docx')
output = source.with_name('烟火哨兵——公共安全智能预警与秒级响应系统_沿用原封面纸张美化版.docx')
document, reference = Document(source), Document(reference_path)
original_text = [paragraph.text for paragraph in document.paragraphs]
original_tables = [[cell.text for row in table.rows for cell in row.cells] for table in document.tables]

def copy_property(destination, original, tag):
    previous = destination.find(qn(tag))
    if previous is not None:
        destination.remove(previous)
    replacement = original.find(qn(tag))
    if replacement is not None:
        destination.insert(0, deepcopy(replacement))

for index in range(6):
    paragraph, model = document.paragraphs[index], reference.paragraphs[index]
    copy_property(paragraph._p, model._p, 'w:pPr')
    for run_index, run in enumerate(paragraph.runs):
        model_run = model.runs[min(run_index, len(model.runs) - 1)]
        copy_property(run._r, model_run._r, 'w:rPr')

cover_table, model_table = document.tables[0], reference.tables[0]
for tag in ['w:tblPr', 'w:tblGrid']:
    copy_property(cover_table._tbl, model_table._tbl, tag)
for row_index, row in enumerate(cover_table.rows):
    for cell_index, cell in enumerate(row.cells):
        model_cell = model_table.rows[row_index].cells[cell_index]
        copy_property(cell._tc, model_cell._tc, 'w:tcPr')
        for paragraph in cell.paragraphs:
            model_paragraph = model_cell.paragraphs[0]
            copy_property(paragraph._p, model_paragraph._p, 'w:pPr')
            for run in paragraph.runs:
                copy_property(run._r, model_paragraph.runs[0]._r, 'w:rPr')

sections = document.element.xpath('.//w:sectPr')
models = reference.element.xpath('.//w:sectPr')
assert len(sections) == len(models) == 3
for index, (section, model) in enumerate(zip(sections, models)):
    copy_property(section, model, 'w:pgMar')
    model_reference = model.find(qn('w:headerReference'))
    model_header = reference.part.related_parts[model_reference.get(qn('r:id'))]
    header = HeaderPart.new(document.part.package)
    header._element = deepcopy(model_header._element)
    for relation in model_header.rels.values():
        if relation.reltype == Relations.IMAGE:
            new_id, image = header.get_or_add_image(BytesIO(relation.target_part.blob))
            for element in header._element.iter():
                for attribute in [qn('r:embed'), qn('r:id')]:
                    if element.get(attribute) == relation.rId:
                        element.set(attribute, new_id)
    for old_reference in list(section.findall(qn('w:headerReference'))):
        section.remove(old_reference)
    new_reference = OxmlElement('w:headerReference')
    new_reference.set(qn('w:type'), 'default')
    new_reference.set(qn('r:id'), document.part.relate_to(header, Relations.HEADER))
    section.insert(0, new_reference)

assert original_text == [paragraph.text for paragraph in document.paragraphs]
assert original_tables == [[cell.text for row in table.rows for cell in row.cells] for table in document.tables]
document.save(output)
with ZipFile(source) as original, ZipFile(output) as polished:
    original_media = Counter(original.read(name) for name in original.namelist() if name.startswith('word/media/') and not name.endswith('/'))
    output_media = Counter(polished.read(name) for name in polished.namelist() if name.startswith('word/media/') and not name.endswith('/'))
    assert not original_media - output_media
print(output)
print('Original content and images preserved; original cover and page backgrounds reused')
