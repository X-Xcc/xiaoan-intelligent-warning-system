from __future__ import annotations

import math
import shutil
from pathlib import Path
from textwrap import wrap
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(r"D:\CICSIC")
OUT = ROOT / "outputs" / "patent_upload_10000565931655"
TITLE = "一种夜间商圈多源事件协同分析与处置流程系统及方法"


def set_font(run, name: str, size: int, bold: bool = False, color: str = "000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def set_margins(section):
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.style = doc.styles["Heading 1"]
        run = p.add_run(text)
        set_font(run, "SimSun", 14, bold=True)
    elif level == 2:
        p.style = doc.styles["Heading 2"]
        run = p.add_run(text)
        set_font(run, "SimSun", 12, bold=True)
    else:
        p.style = doc.styles["Heading 3"]
        run = p.add_run(text)
        set_font(run, "SimSun", 11, bold=True)
    set_paragraph_spacing(p, before=8, after=4, line=1.15)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, "SimSun", 11)
    set_paragraph_spacing(p, before=0, after=6, line=1.2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_claim(doc, text, level=0):
    p = doc.add_paragraph(style="List Number")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 * level)
    run = p.add_run(text)
    set_font(run, "SimSun", 11)
    set_paragraph_spacing(p, before=0, after=4, line=1.15)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_font(run, "SimSun", 11)
    set_paragraph_spacing(p, before=0, after=3, line=1.15)
    return p


def build_doc(path: Path, title: str, body_builder, subtitle: str | None = None):
    doc = Document()
    section = doc.sections[0]
    set_margins(section)

    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            try:
                style.font.name = "SimSun"
                style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
            except Exception:
                pass

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title)
    set_font(run, "SimSun", 16, bold=True)
    set_paragraph_spacing(title_p, before=0, after=6, line=1.0)

    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub.add_run(subtitle)
        set_font(run, "SimSun", 10)
        set_paragraph_spacing(sub, before=0, after=8, line=1.0)

    body_builder(doc)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    footer.runs[0].font.name = "SimSun"
    footer.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    footer.runs[0].font.size = Pt(9)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def patent_description(doc: Document):
    add_heading(doc, "技术领域", 1)
    add_body(doc, "本发明涉及夜间商圈安全治理与事件处置技术领域，尤其涉及一种面向群众上报、视频提示、巡防处置和指挥协同的多源事件协同分析与处置流程系统及方法。")

    add_heading(doc, "背景技术", 1)
    add_body(doc, "夜间商圈、文旅街区和夜游场景通常具有人员流动快、临时摊位多、信息来源分散和处置链路长等特点。现有方式多依赖电话报送、人工巡查和分散台账，存在事件进入慢、派单不准、状态不同步、复盘材料缺失的问题。")
    add_body(doc, "因此，需要一种能够汇聚群众上报、巡防反馈、视频检测和设备告警，并能自动完成事件分级、任务分派、处置回填和审计留痕的技术方案。")

    add_heading(doc, "发明内容", 1)
    add_body(doc, "本发明提供一种夜间商圈多源事件协同分析与处置流程系统及方法，通过事件采集、数据融合、风险分析、任务分派、处置回填和归档复盘六个环节实现流程管理。")
    add_bullet(doc, "事件采集单元用于接入群众端小程序、巡防终端、视频检测结果和人工报警信息。")
    add_bullet(doc, "数据融合单元用于对位置、时间、事件类型、图像证据和历史记录进行统一归并。")
    add_bullet(doc, "风险分析单元用于生成风险等级、优先级和推荐处置建议。")
    add_bullet(doc, "任务分派单元用于按照距离、在线状态、职责范围和任务负载选择处置人员。")
    add_bullet(doc, "处置回填单元用于记录接收、到达、处理、完成和补充说明状态。")
    add_bullet(doc, "归档复盘单元用于沉淀台账、统计高发点位并生成复盘材料。")

    add_heading(doc, "附图说明", 1)
    add_body(doc, "图1为本发明系统结构示意图；图2为本发明事件流程示意图；图3为摘要附图，用于展示系统的主要处理链路。")

    add_heading(doc, "具体实施方式", 1)
    add_body(doc, "以下结合夜间商圈场景对本发明作进一步说明。群众端小程序、巡防人员终端和视频检测模块均接入统一事件接口，系统对新事件进行标准化编码，并按预设规则形成待分析队列。")
    add_body(doc, "当事件满足聚集、纠纷、遗失、求助或异常停留等条件时，风险分析单元输出风险等级和建议处置半径；任务分派单元依据人员位置和在线状态选择最近的巡防人员并推送任务。")
    add_body(doc, "巡防人员在终端完成接收、到场、处置和回填，回填结果进入审计台账。系统同时保留图像证据、时间戳和状态变更记录，以便后续复盘、训练和统计。")
    add_body(doc, "在一种优选实施方式中，事件状态至少包括已提交、已派单、已接收、已到达、处理中、已完成六种状态，且每一次状态变化均记录操作者、时间和位置。")


def patent_claims(doc: Document):
    add_claim(doc, "一种夜间商圈多源事件协同分析与处置流程系统，其特征在于，包括事件采集单元、数据融合单元、风险分析单元、任务分派单元、处置回填单元和归档复盘单元；其中，所述事件采集单元用于接入群众端上报信息、巡防终端信息、视频检测信息和人工报警信息，所述数据融合单元用于对同一事件的多来源信息进行统一编码和归并，所述风险分析单元用于输出事件等级和处置建议，所述任务分派单元用于将事件推送至匹配的巡防人员终端，所述处置回填单元用于记录事件的接收、到达、处理和完成状态，所述归档复盘单元用于形成事件台账和复盘材料。")
    add_claim(doc, "根据权利要求1所述的系统，其特征在于，所述事件采集单元包括群众端小程序接口、巡防终端接口、视频检测接口和设备告警接口。")
    add_claim(doc, "根据权利要求1所述的系统，其特征在于，所述数据融合单元依据事件位置、时间戳、事件类型和证据文件对重复上报进行合并处理。")
    add_claim(doc, "根据权利要求1所述的系统，其特征在于，所述风险分析单元按照预设规则结合事件类型、历史高发程度和现场图像信息生成优先级。")
    add_claim(doc, "根据权利要求1所述的系统，其特征在于，所述任务分派单元依据巡防人员的在线状态、当前位置、职责范围和当前负载选择处置对象。")
    add_claim(doc, "根据权利要求1所述的系统，其特征在于，所述处置回填单元按时间顺序记录已提交、已派单、已接收、已到达、处理中和已完成状态。")
    add_claim(doc, "根据权利要求1所述的系统，其特征在于，所述处置回填单元还用于上传现场图片、语音说明、文字备注和后续跟进标记。")
    add_claim(doc, "根据权利要求1所述的系统，其特征在于，所述归档复盘单元按点位、时段和事件类型生成统计结果与复盘摘要。")
    add_claim(doc, "根据权利要求1所述的系统，其特征在于，所述系统还包括指挥显示单元，所述指挥显示单元用于展示事件队列、在线巡防人员、处置进度和归档结果。")
    add_claim(doc, "一种夜间商圈多源事件协同分析与处置流程方法，其特征在于，包括：接收多源事件信息；对多源事件信息进行统一编码和归并；对事件进行风险分析并生成处置建议；依据巡防人员状态进行任务分派；在巡防终端完成处置回填；以及对事件结果进行归档复盘。")


def patent_abstract(doc: Document):
    add_body(doc, "摘要：本发明公开了一种夜间商圈多源事件协同分析与处置流程系统及方法。该系统接入群众端上报、巡防终端、视频检测和设备告警等多源信息，对同一事件进行统一编码、数据融合和风险分析，并依据巡防人员在线状态、位置和职责范围完成任务分派。巡防人员通过终端完成接收、到达、处理和结果回填，系统同步形成事件台账、审计记录和复盘材料。该方案能够提升夜间商圈事件进入效率、派单准确率和处置完成率，适用于夜游街区、夜市商圈和文旅开放空间的安全治理。")


def create_figures():
    font_path = r"C:\Windows\Fonts\msyh.ttc"
    font_bold_path = r"C:\Windows\Fonts\msyhbd.ttc"
    font = ImageFont.truetype(font_path, 28)
    font_small = ImageFont.truetype(font_path, 24)
    font_bold = ImageFont.truetype(font_bold_path, 30)

    def make_canvas(name: str, subtitle: str):
        img = Image.new("RGB", (1800, 1200), "white")
        d = ImageDraw.Draw(img)
        d.text((60, 30), name, fill="black", font=font_bold)
        d.text((60, 85), subtitle, fill="black", font=font_small)
        return img, d

    def box(d, xy, text, fill=None, outline="black", radius=18):
        d.rounded_rectangle(xy, radius=radius, outline=outline, width=4, fill=fill)
        x0, y0, x1, y1 = xy
        wrapped = "\n".join(wrap(text, width=max(10, int((x1 - x0) / 40))))
        bbox = d.multiline_textbbox((0, 0), wrapped, font=font_small, spacing=8, align="center")
        tw = bbox[2] - bbox[0]
        th = bbox[3] - bbox[1]
        d.multiline_text(((x0 + x1 - tw) / 2, (y0 + y1 - th) / 2), wrapped, fill="black", font=font_small, spacing=8, align="center")

    img, d = make_canvas("图1  系统结构示意图", "多源输入 -> 统一编码 -> 风险分析 -> 任务分派 -> 处置回填 -> 归档复盘")
    left_x = 90
    top = 220
    w = 280
    h = 120
    gap = 40
    inputs = ["群众端上报", "巡防终端", "视频检测", "设备告警"]
    for i, t in enumerate(inputs):
        box(d, (left_x, top + i * (h + gap), left_x + w, top + i * (h + gap) + h), t, fill="#F3F6FA")
    box(d, (520, 330, 780, 510), "事件融合与编码", fill="#EEF4FF")
    box(d, (900, 330, 1180, 510), "风险分析与优先级", fill="#EEF4FF")
    box(d, (1300, 330, 1580, 510), "任务分派与跟踪", fill="#EEF4FF")
    box(d, (900, 720, 1180, 900), "处置回填", fill="#F5F8F2")
    box(d, (1300, 720, 1580, 900), "归档复盘", fill="#F5F8F2")
    for i in range(4):
        y = top + i * (h + gap) + h / 2
        d.line((left_x + w, y, 520, 420), fill="black", width=4)
    d.line((780, 420, 900, 420), fill="black", width=4)
    d.line((1180, 420, 1300, 420), fill="black", width=4)
    d.line((1440, 510, 1440, 720), fill="black", width=4)
    d.line((1180, 810, 1300, 810), fill="black", width=4)
    d.line((1440, 900, 1440, 1040), fill="black", width=4)
    d.line((1180, 810, 900, 810), fill="black", width=4)
    d.text((810, 1040), "状态、证据、时效与审计记录进入统一台账。", fill="black", font=font_small)
    img.save(OUT / "图1-系统结构示意图.png")

    img2, d2 = make_canvas("图2  流程示意图", "以事件状态流转为主线")
    steps = ["事件提交", "统一编码", "风险分析", "任务派单", "到场处置", "结果回填", "归档复盘"]
    x = 100
    y = 350
    step_w = 210
    for i, t in enumerate(steps):
        box(d2, (x + i * 230, y, x + i * 230 + step_w, y + 130), t, fill="#FAFAFA")
        if i < len(steps) - 1:
            d2.line((x + i * 230 + step_w, y + 65, x + (i + 1) * 230, y + 65), fill="black", width=4)
            d2.polygon([(x + i * 230 + step_w + 10, y + 65), (x + i * 230 + step_w - 18, y + 50), (x + i * 230 + step_w - 18, y + 80)], fill="black")
    d2.text((100, 620), "每一次状态变化均记录时间、位置、操作者和证据文件。", fill="black", font=font_small)
    img2.save(OUT / "图2-流程示意图.png")

    img3, d3 = make_canvas("图3  摘要附图", "摘要中建议使用的主链路示意")
    box(d3, (170, 420, 500, 600), "多源采集", fill="#F3F6FA")
    box(d3, (620, 420, 960, 600), "分析与派单", fill="#EEF4FF")
    box(d3, (1080, 420, 1420, 600), "处置与回填", fill="#F5F8F2")
    box(d3, (1490, 420, 1690, 600), "归档", fill="#FAFAFA")
    d3.line((500, 510, 620, 510), fill="black", width=4)
    d3.line((960, 510, 1080, 510), fill="black", width=4)
    d3.line((1420, 510, 1490, 510), fill="black", width=4)
    d3.text((180, 710), "用于表达夜间商圈事件从采集到流程的核心链路。", fill="black", font=font_small)
    img3.save(OUT / "摘要附图.png")


def build_upload_note():
    doc = Document()
    sec = doc.sections[0]
    set_margins(sec)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("专利上传清单")
    set_font(r, "SimSun", 16, bold=True)
    add_body(doc, "1. 说明书：申请发明专利-说明书.docx / PDF")
    add_body(doc, "2. 权利要求书：申请发明专利-权利要求书.docx / PDF")
    add_body(doc, "3. 说明书摘要：申请发明专利-摘要.docx / PDF")
    add_body(doc, "4. 说明书附图：图1-系统结构示意图.png、图2-流程示意图.png")
    add_body(doc, "5. 摘要附图：摘要附图.png")
    add_body(doc, "6. 请求书和申请人信息需在国知局表单内补全；本包不代填身份信息。")
    doc.save(OUT / "上传清单.docx")


def export_pdf_from_docx(docx_path: Path, pdf_path: Path):
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    doc = Document(docx_path)
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "base",
        parent=styles["Normal"],
        fontName="STSong-Light",
        fontSize=11,
        leading=15,
        spaceAfter=6,
    )
    title = ParagraphStyle(
        "title",
        parent=base,
        fontSize=16,
        leading=20,
        alignment=1,
        spaceAfter=8,
    )
    subtitle = ParagraphStyle(
        "subtitle",
        parent=base,
        fontSize=10,
        leading=13,
        alignment=1,
        spaceAfter=8,
    )
    h1 = ParagraphStyle(
        "h1",
        parent=base,
        fontSize=14,
        leading=18,
        spaceBefore=8,
        spaceAfter=4,
        leftIndent=0,
    )
    h2 = ParagraphStyle(
        "h2",
        parent=base,
        fontSize=12,
        leading=16,
        spaceBefore=8,
        spaceAfter=4,
    )
    bullet = ParagraphStyle(
        "bullet",
        parent=base,
        leftIndent=18,
        firstLineIndent=-12,
        spaceAfter=4,
    )
    body = ParagraphStyle("body", parent=base)
    number = ParagraphStyle(
        "number",
        parent=base,
        leftIndent=18,
        firstLineIndent=0,
        spaceAfter=4,
    )

    story = []
    claim_counter = 0
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        style_name = p.style.name
        if idx == 0:
            story.append(Paragraph(text, title))
        elif idx == 1 and p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            story.append(Paragraph(text, subtitle))
        elif style_name == "Heading 1":
            story.append(Paragraph(text, h1))
        elif style_name == "Heading 2":
            story.append(Paragraph(text, h2))
        elif style_name == "Heading 3":
            story.append(Paragraph(text, h2))
        elif style_name == "List Bullet":
            story.append(Paragraph(f"• {text}", bullet))
        elif style_name == "List Number":
            claim_counter += 1
            story.append(Paragraph(f"{claim_counter}. {text}", number))
        else:
            story.append(Paragraph(text, body))
    doc_template = SimpleDocTemplate(
        str(pdf_path),
        pagesize=letter,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )

    def footer(canvas, doc):
        canvas.setFont("STSong-Light", 9)
        canvas.drawCentredString(letter[0] / 2.0, 0.55 * inch, str(canvas.getPageNumber()))

    doc_template.build(story, onFirstPage=footer, onLaterPages=footer)


def zip_package():
    zip_path = OUT / "专利上传件_10000565931655.zip"
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        for p in sorted(OUT.iterdir()):
            if p.is_file() and p.suffix.lower() in {".docx", ".pdf", ".png"}:
                zf.write(p, arcname=p.name)
    return zip_path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    description_docx = OUT / "申请发明专利-说明书.docx"
    claims_docx = OUT / "申请发明专利-权利要求书.docx"
    abstract_docx = OUT / "申请发明专利-摘要.docx"
    upload_docx = OUT / "上传清单.docx"
    build_doc(description_docx, TITLE, patent_description, "夜间商圈安全治理场景")
    build_doc(claims_docx, TITLE, patent_claims, "共10项权利要求")
    build_doc(abstract_docx, TITLE, patent_abstract, "摘要附图见单独文件")
    build_upload_note()
    create_figures()
    export_pdf_from_docx(description_docx, OUT / "申请发明专利-说明书.pdf")
    export_pdf_from_docx(claims_docx, OUT / "申请发明专利-权利要求书.pdf")
    export_pdf_from_docx(abstract_docx, OUT / "申请发明专利-摘要.pdf")
    export_pdf_from_docx(upload_docx, OUT / "上传清单.pdf")
    zip_package()


if __name__ == "__main__":
    main()

