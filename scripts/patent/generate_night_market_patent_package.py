from __future__ import annotations

import json
import shutil
from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\CICSIC")
OUT = ROOT / "outputs" / "patent_application_yanhuo_shaobing_20260902"
REVIEW = OUT / "human_review"
FIGURES = OUT / "附图"

TITLE = "一种面向夜间商圈的多源安全事件统一入库、证据帧复核与就近派单闭环方法及系统"
SHORT_TITLE = "夜间商圈多源安全事件证据帧复核与闭环派单"

SOURCE_ANSWERS = [
    r"D:\CICSIC\README.md",
    r"D:\CICSIC\docs\烟火哨兵交付说明.md",
    r"D:\CICSIC\docs\烟火哨兵项目计划书.md",
    r"D:\CICSIC\docs\烟火哨兵项目实现文档.md",
    r"D:\CICSIC\server\app\services\security_detection.py",
    r"D:\CICSIC\server\app\services\security_ai.py",
    r"D:\CICSIC\server\app\services\event_store.py",
    r"D:\CICSIC\server\app\services\yolo_bridge.py",
    r"D:\CICSIC\server\app\api\routes\security_ai.py",
    r"D:\CICSIC\server\app\api\routes\security_video.py",
]

CLAIM_SUPPORT_ROWS = [
    ["1", "视频源侧帧采样、行为检测、证据帧保存、异步上报、统一事件归并、画面复核、风险更新、就近派单、状态同步", "README.md, security_ai.py, yolo_bridge.py, event_store.py"],
    ["2-4", "证据帧先存储、只传结构化结果、eventKey 幂等归并", "yolo_bridge.py, security_ai.py, event_store.py"],
    ["5-7", "复核服务返回结构化风险结果，失败时转本地待确认", "security_ai.py, dashboard review code"],
    ["8-10", "事件位置与巡防人员位置/在线/职责/负载联动派单，状态回写", "event_store.py, miniprogram workbench"],
    ["11-13", "系统、设备、介质三类权利要求的装置化表达", "claim architecture"],
]

OPEN_ITEMS = [
    "申请人名称、统一社会信用代码、地址、联系人、电话、邮箱待填写。",
    "发明人名单按实际贡献和权属关系最终确认。",
    "如申请前已公开展示、试用、上线或销售，应先核对公开日期与权利影响。",
    "如需要保留具体模型名、阈值或地图服务接口，请与代理人确认后再收窄。",
]


def set_run_font(run, name: str = "SimSun", size: int = 11, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_spacing(paragraph, before: int = 0, after: int = 6, line: float = 1.15) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:ascii"), "SimSun")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "SimSun")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_title(doc: Document, title: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=5, line=1.0)
    run = p.add_run(title)
    set_run_font(run, size=16, bold=True)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p2, after=10, line=1.0)
        run2 = p2.add_run(subtitle)
        set_run_font(run2, size=10)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    set_spacing(p, before=10 if level == 1 else 7, after=4, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size={1: 14, 2: 12, 3: 11}[level], bold=True)


def add_body(doc: Document, text: str, indent: bool = True) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.28) if indent else Inches(0)
    set_spacing(p, after=6, line=1.15)
    run = p.add_run(text)
    set_run_font(run)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_spacing(p, after=3, line=1.15)
    run = p.add_run(text)
    set_run_font(run)


def add_claim(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_spacing(p, after=5, line=1.12)
    run = p.add_run(text)
    set_run_font(run)


def finish_doc(doc: Document, path: Path) -> None:
    footer = doc.sections[0].footer.paragraphs[0]
    add_page_number(footer)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row_values):
            cell = cells[idx]
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def md_table(headers: list[str], rows: list[list[str]]) -> str:
    head = "| " + " | ".join(headers) + " |"
    sep = "| " + " | ".join(["---"] * len(headers)) + " |"
    body = "\n".join("| " + " | ".join(row) + " |" for row in rows)
    return "\n".join([head, sep, body])


def write_markdown_case(path: Path, title: str, subtitle: str, blocks: list[dict]) -> None:
    lines = [f"# {title}", f"**{subtitle}**", ""]
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            lines.append(f"{'#' * block.get('level', 2)} {block['text']}")
            lines.append("")
        elif kind == "para":
            lines.append(block["text"])
            lines.append("")
        elif kind == "bullets":
            for item in block["items"]:
                lines.append(f"- {item}")
            lines.append("")
        elif kind == "table":
            lines.append(md_table(block["headers"], block["rows"]))
            lines.append("")
        elif kind == "note":
            lines.append(f"> {block['text']}")
            lines.append("")
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def build_block_docx(path: Path, title: str, subtitle: str, blocks: list[dict]) -> None:
    doc = Document()
    configure_doc(doc)
    add_title(doc, title, subtitle)
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            add_heading(doc, block["text"], block.get("level", 1))
        elif kind == "para":
            add_body(doc, block["text"], indent=block.get("indent", True))
        elif kind == "bullets":
            for item in block["items"]:
                add_bullet(doc, item)
        elif kind == "table":
            add_table(doc, block["headers"], block["rows"])
        elif kind == "note":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)
            set_spacing(p, after=6, line=1.15)
            run = p.add_run(block["text"])
            set_run_font(run)
        elif kind == "image":
            if block.get("page_break_before"):
                doc.add_page_break()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            p.add_run().add_picture(str(block["path"]), width=Inches(block.get("width", 6.5)))
            if block.get("caption"):
                c = doc.add_paragraph()
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_spacing(c, after=8, line=1.0)
                run = c.add_run(block["caption"])
                set_run_font(run, size=10, bold=True)
    finish_doc(doc, path)


def build_description() -> None:
    doc = Document()
    configure_doc(doc)
    add_title(doc, TITLE, "说明书")

    add_heading(doc, "技术领域")
    add_body(
        doc,
        "本发明涉及视频智能分析、事件数据处理、移动通信和安全指挥调度技术领域，尤其涉及一种面向夜间商圈的多源安全事件统一入库、证据帧复核与就近派单闭环方法及系统。",
    )

    add_heading(doc, "背景技术")
    add_body(
        doc,
        "夜间商圈、夜游街区和开放式文旅空间具有人员流动性强、摄像点位分散、突发事件来源多样以及巡防力量动态变化等特点。现有安全管理系统通常将视频监控、群众求助、巡防记录和指挥调度分别部署，视频检测结果与人工上报难以形成统一事件记录。",
    )
    add_body(
        doc,
        "在一种常见处理方式中，视频检测端持续向平台传输视频或只发送简单告警文本。前一种方式带来带宽、隐私和平台处理压力，后一种方式缺少可复核的现场证据；同时，告警去重、人员选择、路线推荐和处置状态又往往依靠人工操作，导致同一事件重复建单、派单对象不匹配以及处置过程缺少可审计证据。",
    )
    add_body(
        doc,
        "因此，需要一种能够在检测端快速提取事件线索和证据帧，在保存证据后异步发送结构化结果，并在统一事件平台中完成复核、风险更新、就近派单、状态同步和审计归档的技术方案。",
    )

    add_heading(doc, "发明内容")
    add_heading(doc, "发明目的", 2)
    add_body(
        doc,
        "本发明的目的在于提供一种面向夜间商圈的多源安全事件统一入库、证据帧复核与就近派单闭环方法及系统，以解决多源事件难以归并、视频告警缺乏证据、复核过程与事件台账脱节以及派单和处置状态不同步的问题。",
    )

    add_heading(doc, "技术方案", 2)
    add_body(
        doc,
        "为实现上述目的，本发明采用如下技术方案。系统包括视频采集模块、快速检测模块、证据帧存储模块、结果桥接模块、事件接入模块、事件归并模块、视觉复核模块、风险处理模块、派单路线模块、状态同步模块和审计归档模块。",
    )
    add_bullet(doc, "视频采集模块从固定摄像头、移动摄像头或其他视频源获取连续视频数据。")
    add_bullet(doc, "快速检测模块在视频源侧对采样帧执行目标行为检测，输出动作标签、人员数量、相机标识、时间戳、帧数、帧率和证据帧文件名。")
    add_bullet(doc, "证据帧存储模块先将证据帧和检测记录写入检测端的本地存储或共享存储，再生成与检测记录对应的证据索引。")
    add_bullet(doc, "结果桥接模块根据检测记录生成结构化上报载荷，通过异步接口向事件平台发送动作标签、检测统计、相机信息、事件标识和证据帧或证据地址，而不发送连续视频流。")
    add_bullet(doc, "事件接入模块接收视频检测、群众求助、商户上报、巡防反馈和人工报警中的至少两类数据，并将其转换为统一事件数据结构。")
    add_bullet(doc, "事件归并模块依据来源标识、事件标识、相机标识、位置、时间窗口和动作标签进行幂等判断，对重复上报进行合并或关联。")
    add_bullet(doc, "视觉复核模块将证据帧和结构化检测信息发送给画面复核服务，获得是否存在目标风险、风险等级、人数估计、场景摘要和处置建议等结构化复核结果。")
    add_bullet(doc, "风险处理模块根据检测结果和复核结果更新事件风险等级、优先级、是否需要派单以及待人工确认状态。")
    add_bullet(doc, "派单路线模块依据事件位置、巡防人员当前位置、在线状态、职责范围和当前任务负载选择处置人员，并计算或调用对应的响应路线和预计到达时间。")
    add_bullet(doc, "状态同步模块将事件的提交、派单、接收、到达、处理中和完成状态同步至指挥端、巡防端和必要的群众端。")
    add_bullet(doc, "审计归档模块记录事件来源、证据索引、风险变化、操作人员、时间、位置、状态变化、派单结果和处置结果，形成可追溯的事件记录。")

    add_heading(doc, "进一步技术方案", 2)
    add_body(
        doc,
        "优选地，快速检测模块仅对满足预设行为条件的检测结果生成证据帧，例如人员聚集、异常聚集、打架、跌倒或值守离岗；对不满足条件的帧不生成上报载荷，以降低存储和网络传输量。",
    )
    add_body(
        doc,
        "优选地，统一事件数据结构至少包括 sourceId、eventKey、actions、primaryAction、timestamp、personCount、frameCount、fps、cameraId、cameraName、imageFilename、sourcePath 和 eventId 中的多项字段。事件平台在写入事件前根据 eventKey 执行幂等检查，使同一检测结果的重复发送不会生成多个事件工单。",
    )
    add_body(
        doc,
        "优选地，视觉复核服务按照预设格式返回 JSON 对象，并由视觉复核模块执行字段校验和标准化，将风险等级限定在预设集合内；当复核服务不可用或未配置时，根据本地检测动作和人数生成待人工确认的本地复核结果，保证事件仍可进入人工处置流程。",
    )
    add_body(
        doc,
        "优选地，派单路线模块先根据事件所在摄像点或上报位置确定事件坐标，再从在线巡防人员集合中筛选职责匹配的人员，按照位置距离、预计到达时间和任务负载计算排序结果，将排序第一的人员作为推荐处置对象，并将路线信息写入事件记录。",
    )

    add_heading(doc, "有益效果")
    add_bullet(doc, "通过证据帧先保存、检测结果后异步上报，使视频检测端与事件平台解耦，降低连续视频传输带来的带宽和隐私压力。")
    add_bullet(doc, "通过统一事件结构和幂等事件标识，将视频检测、群众求助、商户上报和巡防反馈纳入同一事件链路，减少重复建单。")
    add_bullet(doc, "通过结构化画面复核结果更新风险等级和处置建议，使视频告警具有可追溯的复核证据。")
    add_bullet(doc, "通过结合位置、在线状态、职责和任务负载进行人员选择，使派单结果与现场响应条件相匹配。")
    add_bullet(doc, "通过状态同步和审计归档保留从发现到完成的全过程数据，为指挥调度、复盘分析和训练使用提供统一数据基础。")

    add_heading(doc, "附图说明")
    add_body(doc, "图1为本发明系统结构示意图。")
    add_body(doc, "图2为本发明事件处理方法流程示意图。")
    add_body(doc, "图3为本发明证据帧复核与风险更新流程示意图。")
    add_body(doc, "图4为本发明事件状态及审计记录示意图。")

    add_heading(doc, "具体实施方式")
    add_body(
        doc,
        "下面结合附图和夜间商圈安全治理场景对本发明作进一步说明。图1中，视频采集模块100连接快速检测模块110，快速检测模块110连接证据帧存储模块120和结果桥接模块130。群众端、商户端、巡防端和人工报警终端分别通过对应接口连接事件接入模块210。结果桥接模块130通过异步网络接口连接事件接入模块210。",
    )
    add_body(
        doc,
        "在视频检测过程中，快速检测模块110从连续视频中按照采样周期获取图像帧，对图像帧执行人员目标检测、目标数量统计和行为动作判断。当动作标签属于预设事件动作集合时，快速检测模块110将当前帧或经过压缩处理的关键帧交给证据帧存储模块120保存，并生成检测记录。检测记录包括检测源标识、相机标识、事件键、动作列表、主动作、时间戳、人员数量、帧数、帧率和证据帧索引。",
    )
    add_body(
        doc,
        "证据帧存储模块120完成写入后，结果桥接模块130才创建上报任务。上报任务携带结构化检测记录和证据帧数据，或者携带能够访问证据帧的地址。结果桥接模块130不向事件平台转发视频采集模块100产生的连续视频流。上报任务失败时，检测端保留已保存的证据帧和待发送记录，并根据重试策略再次发送。",
    )
    add_body(
        doc,
        "事件接入模块210接收到检测结果后，事件归并模块220根据 eventKey 查询已有检测记录和事件记录。若存在相同 eventKey，则更新检测记录的时间、动作、人员数量、证据索引和复核状态，不重复创建事件；若不存在，则创建新的视频提示事件，并把检测记录写入安全检测表，把事件写入安全事件表，同时创建指挥中心报警推送记录。",
    )
    add_body(
        doc,
        "视觉复核模块230从检测记录读取证据帧，构造包含相机信息和本地检测信息的复核载荷，并发送至画面复核服务。画面复核服务返回 isGathering、riskLevel、peopleEstimate、sceneSummary 和 suggestion 等字段。结果标准化模块240对返回内容进行 JSON 解析、字段校验和风险等级归一化，并将结果与 eventId 关联写回事件记录。",
    )
    add_body(
        doc,
        "风险处理模块250综合本地检测动作和视觉复核结果确定事件优先级。对于复核确认的人员聚集、通道拥堵或打架斗殴事件，风险处理模块250将事件置于待派单队列；对于复核结果不确定的事件，保留证据帧并标记为待人工确认。即使画面复核服务暂时不可用，也根据本地动作标签和人员数量生成本地复核上下文，使事件能够继续接受人工处理。",
    )
    add_body(
        doc,
        "派单路线模块260读取事件位置和巡防人员数据。事件位置可以由群众上报的经纬度、摄像点预置坐标或网格位置确定。巡防人员数据至少包括人员标识、职责角色、在线状态、当前位置和已有任务数量。派单路线模块260筛选在线且职责匹配的人员，对候选人员计算距离或预计到达时间，并结合当前任务负载形成排序，写入推荐人员和路线信息。",
    )
    add_body(
        doc,
        "状态同步模块270通过事件接口和实时消息通道将事件状态发送至指挥端和巡防端。事件状态按照已提交、已派单、已接收、已到达、处理中和已完成的顺序流转；每一次状态变化均由审计归档模块280记录前一状态、后一状态、操作人员、操作时间、位置、处置结果和证据索引。完成状态写入后，系统生成事件复盘摘要或统计记录。",
    )
    add_body(
        doc,
        "在一具体实施例中，夜市商圈摄像头拍摄到人员聚集，快速检测模块110统计得到人员数量为七人，并判断动作标签为人员聚集。证据帧存储模块120在本地保存 JPEG 证据帧，结果桥接模块130将动作标签、人员数量、相机标识、时间戳和证据帧发送给事件接入模块210。事件归并模块220生成唯一视频事件，视觉复核模块230确认现场存在聚集并输出中风险等级，派单路线模块260根据事件点位和在线巡防人员位置推荐最近的巡防人员。巡防人员完成接收、到达和处置后，状态同步模块270将结果发送给指挥端，审计归档模块280保存全过程记录。",
    )
    add_body(
        doc,
        "在另一具体实施例中，群众端同时提交同一位置的求助信息。事件归并模块220按照位置、时间窗口和事件类型建立关联关系，将群众求助作为同一风险事件的补充来源，并把群众描述与视频检测证据帧共同提供给视觉复核模块230和指挥端。由此形成包含人工描述、检测记录、视觉复核结果和处置日志的统一事件包。",
    )
    add_body(
        doc,
        "本发明还可以由计算机设备执行，所述计算机设备包括处理器、存储器、通信接口和显示接口；处理器执行存储器中的程序指令以实现上述方法。所述程序指令也可以存储在非暂态计算机可读存储介质中。",
    )

    finish_doc(doc, OUT / "说明书.docx")


def build_claims() -> None:
    doc = Document()
    configure_doc(doc)
    add_title(doc, TITLE, "权利要求书")

    claims = [
        "一种面向夜间商圈的多源安全事件统一入库、证据帧复核与就近派单闭环方法，其特征在于，包括：从至少一个视频源获取连续视频数据；对所述连续视频数据进行帧采样和行为检测，得到包含动作标签、人员数量、相机标识、时间戳和检测统计信息的检测记录；当所述动作标签属于预设事件动作集合时，保存与所述检测记录对应的证据帧并生成证据索引；在所述证据帧保存完成后，将所述检测记录以及所述证据帧或所述证据索引异步发送至事件平台；由所述事件平台将视频检测、群众求助、商户上报、巡防反馈和人工报警中的至少两类数据归一为统一事件数据结构，并根据事件标识执行幂等归并；将所述证据帧和所述检测记录发送至画面复核服务，获取结构化复核结果；根据所述检测记录和所述结构化复核结果更新事件风险等级和处置优先级；根据事件位置、巡防人员当前位置、在线状态、职责范围和任务负载选择处置人员并生成响应路线；将事件状态同步至指挥端和巡防端，并记录事件来源、证据索引、风险变化、派单结果、状态变化和处置结果。",
        "根据权利要求1所述的方法，其特征在于，所述异步发送仅发送动作标签、人员数量、帧数、帧率、相机标识、时间戳、事件标识以及证据帧或证据索引，不发送所述视频源产生的连续视频流。",
        "根据权利要求1所述的方法，其特征在于，所述证据帧的保存包括：按照预设图像格式写入本地存储或共享存储；生成包含证据帧文件名、源路径和生成时间的证据索引；在证据帧写入成功后创建上报任务；当上报任务失败时保留检测记录和证据索引并执行重试。",
        "根据权利要求1所述的方法，其特征在于，所述统一事件数据结构至少包括 sourceId、eventKey、actions、primaryAction、timestamp、personCount、frameCount、fps、cameraId、cameraName、imageFilename、sourcePath 和 eventId 中的多项字段。",
        "根据权利要求1所述的方法，其特征在于，所述幂等归并包括：查询与所述事件标识对应的检测记录和事件记录；当查询到已存在的检测记录或事件记录时，更新检测时间、动作标签、人员数量、证据索引或复核状态中的至少一项而不新增事件工单；当未查询到对应记录时，创建视频检测事件并生成报警推送记录。",
        "根据权利要求1所述的方法，其特征在于，所述结构化复核结果至少包括风险是否存在、风险等级、人数估计、场景摘要和处置建议；对所述结构化复核结果执行字段校验和风险等级归一化，并将归一化结果与统一事件数据结构中的事件标识关联存储。",
        "根据权利要求1所述的方法，其特征在于，当画面复核服务不可用或未配置时，根据动作标签和人员数量生成本地复核结果，并将所述事件标记为待人工确认，以使事件继续进入人工派单或人工复核流程。",
        "根据权利要求1所述的方法，其特征在于，所述选择处置人员包括：从巡防人员集合中筛选在线状态有效且职责范围匹配的候选人员；分别计算事件位置与候选人员当前位置之间的距离或预计到达时间；结合候选人员的当前任务负载生成排序；将排序结果中的目标人员、距离、预计到达时间和路线写入事件记录。",
        "根据权利要求1所述的方法，其特征在于，所述事件状态至少包括已提交、已派单、已接收、已到达、处理中和已完成；每次状态变化均记录前一状态、后一状态、操作人员、操作时间、操作位置、处置结果和证据索引。",
        "根据权利要求1所述的方法，其特征在于，所述统一事件数据结构还接收群众上报的地理位置和文字描述，并依据所述地理位置、预设时间窗口和事件类型将群众上报与视频检测事件建立关联关系，形成包含人工描述、视频检测记录和证据帧的事件包。",
        "一种面向夜间商圈的多源安全事件统一入库、证据帧复核与就近派单闭环系统，其特征在于，包括：视频采集模块，用于获取连续视频数据；快速检测模块，用于对采样帧执行行为检测并输出检测记录；证据帧存储模块，用于在上报前保存证据帧并生成证据索引；结果桥接模块，用于在证据帧保存完成后异步发送结构化检测结果；事件接入模块，用于接收视频检测、群众求助、商户上报、巡防反馈和人工报警中的至少两类数据；事件归并模块，用于将不同来源数据归一为统一事件数据结构并执行幂等归并；视觉复核模块，用于基于证据帧和检测记录获取结构化复核结果；风险处理模块，用于更新事件风险等级和处置优先级；派单路线模块，用于依据事件位置、巡防人员当前位置、在线状态、职责范围和任务负载选择处置人员并生成响应路线；状态同步模块，用于同步事件状态；审计归档模块，用于记录事件全过程数据。",
        "一种计算机设备，包括处理器和存储器，其特征在于，所述存储器中存储有程序，所述程序被所述处理器执行时，使所述计算机设备执行权利要求1至10任一项所述的方法。",
        "一种非暂态计算机可读存储介质，其上存储有程序，所述程序被处理器执行时，使所述处理器执行权利要求1至10任一项所述的方法。",
    ]
    for claim in claims:
        add_claim(doc, claim)
    finish_doc(doc, OUT / "权利要求书.docx")


def build_abstract() -> None:
    doc = Document()
    configure_doc(doc)
    add_title(doc, TITLE, "说明书摘要")
    add_body(
        doc,
        "本发明公开一种面向夜间商圈的多源安全事件统一入库、证据帧复核与就近派单闭环方法及系统。视频源侧对采样帧进行行为检测，在生成事件动作、人员数量和相机信息后先保存证据帧，再异步发送结构化检测结果及证据索引。事件平台将视频检测、群众求助、商户上报和巡防反馈归一为统一事件记录并执行幂等归并，基于证据帧调用画面复核服务，获得风险等级、人数估计、场景摘要和处置建议；再结合事件位置、巡防人员位置、在线状态、职责范围和任务负载完成就近派单，并同步提交、派单、接收、到达、处理和完成状态，记录全过程审计信息。本发明能够降低连续视频传输压力，减少重复建单，提高告警复核和现场响应的可追溯性。",
        indent=False,
    )
    p = doc.add_paragraph()
    set_spacing(p, before=12, after=4, line=1.15)
    run = p.add_run("摘要附图：图1")
    set_run_font(run, size=11)
    finish_doc(doc, OUT / "摘要.docx")


def font_pair() -> tuple[ImageFont.FreeTypeFont, ImageFont.FreeTypeFont]:
    regular = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    if not regular.exists():
        regular = Path(r"C:\Windows\Fonts\simsun.ttc")
    if not bold.exists():
        bold = regular
    return ImageFont.truetype(str(regular), 26), ImageFont.truetype(str(bold), 30)


def make_canvas(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw, ImageFont.FreeTypeFont, ImageFont.FreeTypeFont]:
    regular, bold = font_pair()
    image = Image.new("L", (2000, 1300), 255)
    draw = ImageDraw.Draw(image)
    draw.text((70, 36), title, fill=0, font=bold)
    draw.text((70, 94), subtitle, fill=0, font=regular)
    return image, draw, regular, bold


def draw_box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], text: str, font: ImageFont.FreeTypeFont) -> None:
    # Patent application figures use monochrome line drawings.
    draw.rounded_rectangle(xy, radius=16, outline=0, width=4, fill=255)
    x0, y0, x1, y1 = xy
    lines = "\n".join(wrap(text, width=max(8, int((x1 - x0) / 42))))
    bbox = draw.multiline_textbbox((0, 0), lines, font=font, spacing=8, align="center")
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    draw.multiline_text(((x0 + x1 - tw) / 2, (y0 + y1 - th) / 2), lines, fill="black", font=font, spacing=8, align="center")


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line((*start, *end), fill=0, width=4)
    x0, y0 = start
    x1, y1 = end
    if abs(x1 - x0) >= abs(y1 - y0):
        direction = 1 if x1 >= x0 else -1
        draw.polygon([(x1, y1), (x1 - 18 * direction, y1 - 12), (x1 - 18 * direction, y1 + 12)], fill=0)
    else:
        direction = 1 if y1 >= y0 else -1
        draw.polygon([(x1, y1), (x1 - 12, y1 - 18 * direction), (x1 + 12, y1 - 18 * direction)], fill=0)


def elbow_arrow(draw: ImageDraw.ImageDraw, points: list[tuple[int, int]]) -> None:
    for start, end in zip(points, points[1:-1]):
        draw.line((*start, *end), fill=0, width=4)
    arrow(draw, points[-2], points[-1])


def create_figures() -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    image, draw, regular, bold = make_canvas(
        "图1  系统结构示意图",
        "视频源侧证据帧先保存，结构化结果异步进入统一事件平台",
    )
    draw_box(draw, (80, 250, 390, 390), "固定摄像头\n移动摄像头\n其他视频源", regular)
    draw_box(draw, (80, 500, 390, 640), "群众端/商户端\n求助与上报", regular)
    draw_box(draw, (80, 750, 390, 890), "巡防端/人工报警\n现场反馈", regular)
    draw_box(draw, (520, 300, 850, 450), "快速检测\n帧采样与动作判断", regular)
    draw_box(draw, (520, 550, 850, 700), "证据帧存储\n检测记录与索引", regular)
    draw_box(draw, (520, 800, 850, 950), "结果桥接\n异步上报", regular)
    draw_box(draw, (1020, 400, 1370, 550), "事件接入与归并\n统一事件标识", regular)
    draw_box(draw, (1500, 250, 1880, 400), "画面复核服务\n结构化风险结果", regular)
    draw_box(draw, (1500, 500, 1880, 650), "风险更新与\n优先级处理", regular)
    draw_box(draw, (1500, 750, 1880, 900), "位置/在线/职责/负载\n派单与路线", regular)
    draw_box(draw, (1020, 850, 1370, 1000), "状态同步与\n审计归档", regular)
    arrow(draw, (390, 320), (520, 375))
    elbow_arrow(draw, [(390, 570), (450, 570), (450, 180), (950, 180), (950, 450), (1020, 450)])
    elbow_arrow(draw, [(390, 820), (450, 820), (450, 1070), (950, 1070), (950, 520), (1020, 520)])
    arrow(draw, (685, 450), (685, 550))
    arrow(draw, (685, 700), (685, 800))
    arrow(draw, (850, 875), (1020, 475))
    arrow(draw, (1370, 475), (1500, 575))
    arrow(draw, (1690, 400), (1690, 500))
    arrow(draw, (1690, 650), (1690, 750))
    arrow(draw, (1690, 900), (1370, 925))
    arrow(draw, (1200, 550), (1200, 850))
    draw.text((960, 1090), "证据帧、检测结果、复核结果、派单结果和状态日志形成统一事件链路", fill="black", font=regular)
    image.save(FIGURES / "图1-系统结构示意图.png")

    image, draw, regular, bold = make_canvas(
        "图2  事件处理方法流程示意图",
        "从检测端采样到事件闭环的处理顺序",
    )
    labels = ["获取视频", "采样检测", "保存证据帧", "异步上报", "事件归并", "画面复核", "风险更新", "就近派单", "同步归档"]
    x, y = 70, 460
    box_w, box_h, gap = 175, 150, 28
    for index, label in enumerate(labels):
        left = x + index * (box_w + gap)
        draw_box(draw, (left, y, left + box_w, y + box_h), label, regular)
        if index < len(labels) - 1:
            arrow(draw, (left + box_w, y + box_h // 2), (left + box_w + gap, y + box_h // 2))
    draw.text((120, 760), "证据帧保存成功是异步上报任务进入发送队列的前置条件", fill="black", font=regular)
    draw.text((120, 820), "重复 eventKey 只更新已有记录；复核失败时保留本地检测结果并转人工确认", fill="black", font=regular)
    image.save(FIGURES / "图2-事件处理流程示意图.png")

    image, draw, regular, bold = make_canvas(
        "图3  证据帧复核与风险更新流程示意图",
        "本地检测上下文与证据帧共同构成复核输入",
    )
    draw_box(draw, (120, 420, 470, 610), "本地检测记录\n动作/人数/时间/相机", regular)
    draw_box(draw, (120, 760, 470, 950), "证据帧\n文件或地址", regular)
    draw_box(draw, (700, 540, 1100, 760), "复核载荷\n相机信息 + 检测上下文\n+ 证据帧", regular)
    draw_box(draw, (1350, 300, 1830, 470), "结构化复核结果\n风险等级/人数估计\n场景摘要/处置建议", regular)
    draw_box(draw, (1350, 650, 1830, 820), "事件风险与优先级\n待人工确认或待派单", regular)
    draw_box(draw, (1350, 980, 1830, 1150), "写回事件记录\n关联 eventId 与审计日志", regular)
    arrow(draw, (470, 515), (700, 625))
    arrow(draw, (470, 855), (700, 675))
    arrow(draw, (1100, 650), (1350, 385))
    arrow(draw, (1590, 470), (1590, 650))
    arrow(draw, (1590, 820), (1590, 980))
    image.save(FIGURES / "图3-证据帧复核流程示意图.png")

    image, draw, regular, bold = make_canvas(
        "图4  事件状态及审计记录示意图",
        "状态流转与操作、时间、位置、证据同步记录",
    )
    states = ["已提交", "已派单", "已接收", "已到达", "处理中", "已完成"]
    x, y = 110, 420
    for index, state in enumerate(states):
        left = x + index * 300
        draw_box(draw, (left, y, left + 220, y + 130), state, regular)
        if index < len(states) - 1:
            arrow(draw, (left + 220, y + 65), (left + 300, y + 65))
    draw_box(draw, (500, 760, 1500, 990), "审计记录：事件标识、前一状态、后一状态、操作人员、操作时间、操作位置、风险变化、派单结果、处置结果、证据索引", regular)
    for left in [220, 520, 820, 1120, 1420]:
        arrow(draw, (left, 550), (left + 280, 760))
    image.save(FIGURES / "图4-事件状态审计示意图.png")

    shutil.copyfile(FIGURES / "图1-系统结构示意图.png", OUT / "摘要附图.png")


def build_submission_template() -> None:
    text = f"""# 发明专利申请填报模板

## 申请主题

- 发明名称：{TITLE}
- 申请类型：发明专利
- 技术文件：说明书、权利要求书、说明书摘要、说明书附图

## 请求书中需要填写

- 申请人类型：个人 / 企业 / 其他法人
- 申请人名称：
- 申请人证件号码或统一社会信用代码：
- 申请人地址：
- 邮政编码：
- 联系电话：
- 电子邮箱：
- 发明人1：
- 发明人2：
- 其他发明人：
- 联系人：
- 是否要求优先权：否 / 是（如选是，补充在先申请信息）
- 是否提前公开：按正常公开办理
- 是否同时提出实质审查请求：先不提出
- 是否请求费用减缴：按实际资格办理

## 电子申请文件对应关系

1. “说明书”上传 `说明书.docx` 经国知局电子申请工具转换后的 XML。
2. “权利要求书”上传 `权利要求书.docx` 经国知局电子申请工具转换后的 XML。
3. “说明书摘要”上传 `摘要.docx` 经国知局电子申请工具转换后的 XML。
4. “说明书附图”上传 `附图` 目录内的 PNG 文件。
5. 请求书主体信息在网页表单内填写，不把身份证号、统一社会信用代码写入技术文件。

## 申请前核对

- 发明人姓名与实际研发贡献一致。
- 申请人对技术方案具有申请权。
- 技术文件中没有使用尚未公开的密码、密钥或个人敏感信息。
- 如申请人不是发明人本人，留存劳动关系、委托开发或权利转让等权属材料。
- 提交前确认是否存在在先公开、论文、产品演示或销售行为。
"""
    (OUT / "请求书填报模板.md").write_text(text, encoding="utf-8")


def build_evidence_manifest() -> None:
    manifest = {
        "title": TITLE,
        "shortTitle": SHORT_TITLE,
        "applicationType": "invention",
        "createdAt": "2026-09-02",
        "sourceEvidence": [
            r"D:\CICSIC\README.md",
            r"D:\CICSIC\docs\烟火哨兵交付说明.md",
            r"D:\CICSIC\docs\烟火哨兵项目计划书.md",
            r"D:\CICSIC\docs\烟火哨兵项目实现文档.md",
            r"D:\CICSIC\server\app\services\security_detection.py",
            r"D:\CICSIC\server\app\services\security_ai.py",
            r"D:\CICSIC\server\app\services\event_store.py",
            r"D:\CICSIC\server\app\services\yolo_bridge.py",
            r"D:\CICSIC\server\app\api\routes\security_ai.py",
            r"D:\CICSIC\server\app\api\routes\security_video.py",
        ],
        "deliverables": [
            "说明书.docx",
            "权利要求书.docx",
            "摘要.docx",
            "摘要附图.png",
            "附图/图1-系统结构示意图.png",
            "附图/图2-事件处理流程示意图.png",
            "附图/图3-证据帧复核流程示意图.png",
            "附图/图4-事件状态审计示意图.png",
            "请求书填报模板.md",
            "00_专利草案总览.docx",
            "01_技术交底书.docx",
            "02_权利要求书.docx",
            "03_说明书.docx",
            "04_说明书摘要.docx",
            "05_说明书附图.docx",
            "06_摘要附图.docx",
            "07_初步查新笔记.docx",
            "08_发明专利请求书信息确认表.docx",
            "09_查新与差异化分析.docx",
            "提交文件清单.docx",
        ],
        "portalNote": "当前电子申请正式文件按国知局系统要求转换为 XML；本目录中的 DOCX 是可编辑母稿，附图为图像文件。",
    }
    (OUT / "材料清单.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def build_core_doc_copies() -> None:
    shutil.copyfile(OUT / "权利要求书.docx", OUT / "02_权利要求书.docx")
    shutil.copyfile(OUT / "说明书.docx", OUT / "03_说明书.docx")
    shutil.copyfile(OUT / "摘要.docx", OUT / "04_说明书摘要.docx")


def build_overview_pack() -> None:
    blocks = [
        {"kind": "para", "text": "本总览文件用于在原有烟火哨兵技术包的基础上，确认主申请点、权利要求范围、证据来源和填报边界。当前建议仅围绕一件发明专利展开，主题是“多源安全事件统一入库、证据帧复核与就近派单闭环”。", "indent": False},
        {"kind": "table", "headers": ["项目", "内容", "说明"], "rows": [
            ["发明名称", TITLE, "推荐作为主申请名称"],
            ["申请类型", "发明专利", "方法 + 系统 + 设备 + 介质"],
            ["核心技术问题", "多源事件难归并、证据帧与处置链路脱节、派单和状态不同步", "技术问题而非业务问题"],
            ["核心技术手段", "证据帧先保存、结构化结果异步上报、统一事件归并、画面复核、就近派单、状态回写", "与现有项目代码强相关"],
            ["技术效果", "降低连续视频传输压力、减少重复建单、提升复核和派单可追溯性", "可由结构性效果支撑"],
            ["当前状态", "已形成可编辑母稿与黑白附图", "可继续交代理人审校"],
        ]},
        {"kind": "heading", "text": "发明点判断", "level": 2},
        {"kind": "bullets", "items": [
            "主申请点不是单独的识别模型，而是“检测端证据帧先落地 + 事件平台统一归并 + 视觉复核 + 就近派单 + 状态审计”的闭环组合。",
            "可直接支撑方法权利要求、系统权利要求、设备权利要求和计算机可读存储介质权利要求。",
            "当前文字尽量不写死具体模型名、阈值或地图引擎，避免把保护范围锁得太窄。",
        ]},
        {"kind": "heading", "text": "证据来源", "level": 2},
        {"kind": "bullets", "items": [f"{p}" for p in SOURCE_ANSWERS]},
        {"kind": "heading", "text": "权利要求支撑概览", "level": 2},
        {"kind": "table", "headers": ["层级", "支撑内容", "对应来源"], "rows": CLAIM_SUPPORT_ROWS},
        {"kind": "heading", "text": "待确认事项", "level": 2},
        {"kind": "bullets", "items": OPEN_ITEMS},
    ]
    write_markdown_case(OUT / "00_专利草案总览.md", TITLE, "专利申请总览", blocks)
    build_block_docx(OUT / "00_专利草案总览.docx", TITLE, "专利申请总览", blocks)


def build_technical_disclosure_pack() -> None:
    blocks = [
        {"kind": "para", "text": "以下内容用于技术交底，不是最终提交 CNIPA 的正式说明书。它把烟火哨兵现有实现中最可保护的技术机制先固定下来，供后续权利要求和说明书复用。", "indent": False},
        {"kind": "heading", "text": "技术领域", "level": 2},
        {"kind": "para", "text": "本方案涉及视频智能分析、事件数据处理、移动通信、巡防调度和安全指挥技术，尤其涉及一种面向夜间商圈的多源安全事件统一入库、证据帧复核与就近派单闭环方法及系统。"},
        {"kind": "heading", "text": "背景技术", "level": 2},
        {"kind": "para", "text": "夜间商圈存在客流分散、摄像点位多、上报来源杂和巡防位置动态变化等特点。现有系统常把视频告警、群众求助、商户上报、巡防反馈和人工报警分开处理，导致同一风险被重复建单，证据与工单分离，派单对象和处置状态难以统一审计。"},
        {"kind": "para", "text": "烟火哨兵现有代码已经出现了检测结果写入事件库、证据帧保存、结构化复核、风险分级和巡防同步等能力，但这些能力尚未被整理成一套完整、可审查的专利表达，因此需要用专利写法重新梳理成技术方案。"},
        {"kind": "heading", "text": "技术问题", "level": 2},
        {"kind": "para", "text": "要解决的技术问题是：在多源安全事件输入并发、证据帧需要快速保留、复核服务可能异步返回的条件下，如何让同一事件在统一事件平台内完成幂等归并、结构化复核、就近派单、状态同步和审计归档。"},
        {"kind": "heading", "text": "技术方案", "level": 2},
        {"kind": "bullets", "items": [
            "视频源侧只对采样帧执行检测，在满足预设事件动作时先保存证据帧，再异步发送结构化结果，不传连续视频流。",
            "事件平台把视频检测、群众求助、商户上报、巡防反馈和人工报警中的至少两类数据归一成统一事件数据结构，并用 eventKey 或同类唯一标识做幂等归并。",
            "画面复核模块基于证据帧和本地检测上下文调用复核服务，拿到风险等级、人数估计、场景摘要和处置建议。",
            "派单路线模块综合事件位置、巡防人员当前位置、在线状态、职责范围和任务负载，生成推荐处置人和响应路线。",
            "状态同步模块把提交、派单、接收、到达、处理中、已完成等状态写回指挥端和巡防端，并在审计表中留下全过程痕迹。",
        ]},
        {"kind": "heading", "text": "核心发明点", "level": 2},
        {"kind": "bullets", "items": [
            "证据帧先保存，再异步上报结构化结果。",
            "多源事件统一入库与 eventKey 幂等归并。",
            "复核结果反向影响风险等级、派单优先级和处置路线。",
            "事件状态与审计记录同步回写，形成闭环。",
        ]},
        {"kind": "heading", "text": "支撑性特征", "level": 2},
        {"kind": "bullets", "items": [
            "检测结果至少包含动作标签、人员数量、相机标识、时间戳、帧数、帧率和证据帧路径。",
            "复核服务返回结构化 JSON，返回失败时仍保留本地检测结果供人工确认。",
            "巡防人员筛选同时考虑职责匹配与实时在线状态。",
            "事件记录保存来源、证据索引、风险变化、派单结果、处置结果和状态变化。",
        ]},
        {"kind": "heading", "text": "可选特征", "level": 2},
        {"kind": "bullets", "items": [
            "是否使用特定视觉模型名可作为实现细节，不写入主独立权利要求。",
            "是否接入特定地图服务、路径算法或报警推送渠道，可作为从属实施例。",
            "是否对事件类型设置更细阈值，可作为从属权利要求或实施例参数。",
        ]},
        {"kind": "heading", "text": "图纸说明", "level": 2},
        {"kind": "table", "headers": ["图号", "名称", "作用"], "rows": [
            ["图1", "系统结构示意图", "展示检测端、事件平台、复核和派单闭环"],
            ["图2", "事件处理流程示意图", "展示从获取视频到同步归档的顺序"],
            ["图3", "证据帧复核流程示意图", "展示检测上下文与证据帧如何形成复核载荷"],
            ["图4", "事件状态审计示意图", "展示状态流转和审计记录"],
        ]},
        {"kind": "heading", "text": "证据映射", "level": 2},
        {"kind": "table", "headers": ["权利要求要素", "说明书支撑", "参考编号/文件"], "rows": [
            ["视频源侧采样与检测", "说明书具体实施方式已描述", "100/110"],
            ["证据帧保存与异步上报", "说明书技术方案和实施例已描述", "120/130"],
            ["统一事件归并", "说明书详细描述已描述", "210/220"],
            ["结构化复核", "说明书和图3已描述", "230/240"],
            ["就近派单与状态回写", "说明书和图4已描述", "250/260/270/280"],
        ]},
    ]
    write_markdown_case(OUT / "01_技术交底书.md", TITLE, "技术交底书", blocks)
    build_block_docx(OUT / "01_技术交底书.docx", TITLE, "技术交底书", blocks)


def build_figures_companion_pack() -> None:
    blocks = [
        {"kind": "para", "text": "本文件是说明书附图的伴随说明，图像本体仍以 `附图` 目录中的黑白 PNG 为准。", "indent": False},
        {"kind": "table", "headers": ["图号", "文件名", "说明"], "rows": [
            ["图1", "图1-系统结构示意图.png", "展示视频源、检测、证据存储、事件归并、复核、派单和审计"],
            ["图2", "图2-事件处理流程示意图.png", "展示事件处理顺序"],
            ["图3", "图3-证据帧复核流程示意图.png", "展示复核载荷与风险回写"],
            ["图4", "图4-事件状态审计示意图.png", "展示状态流转与审计记录"],
        ]},
        {"kind": "heading", "text": "参考标号", "level": 2},
        {"kind": "table", "headers": ["标号", "名称"], "rows": [
            ["100", "视频采集与检测子系统"],
            ["110", "快速检测模块"],
            ["120", "证据帧存储模块"],
            ["130", "结果桥接模块"],
            ["210", "事件接入模块"],
            ["220", "事件归并模块"],
            ["230", "视觉复核模块"],
            ["240", "结果标准化模块"],
            ["250", "风险处理模块"],
            ["260", "派单路线模块"],
            ["270", "状态同步模块"],
            ["280", "审计归档模块"],
        ]},
        {"kind": "note", "text": "上传 CNIPA 时，图像本体使用附图目录中的 PNG，本文档作为图纸说明和编号索引保留在申请包中。"},
    ]
    write_markdown_case(OUT / "05_说明书附图.md", TITLE, "说明书附图说明", blocks)
    build_block_docx(
        OUT / "05_说明书附图.docx",
        TITLE,
        "说明书附图说明",
        [
            {"kind": "table", "headers": ["图号", "文件名", "说明"], "rows": [
                ["图1", "图1-系统结构示意图.png", "展示视频源、检测、证据存储、事件归并、复核、派单和审计"],
                ["图2", "图2-事件处理流程示意图.png", "展示事件处理顺序"],
                ["图3", "图3-证据帧复核流程示意图.png", "展示复核载荷与风险回写"],
                ["图4", "图4-事件状态审计示意图.png", "展示状态流转与审计记录"],
            ]},
            {"kind": "heading", "text": "参考标号", "level": 2},
            {"kind": "table", "headers": ["标号", "名称"], "rows": [
                ["100", "视频采集与检测子系统"],
                ["110", "快速检测模块"],
                ["120", "证据帧存储模块"],
                ["130", "结果桥接模块"],
                ["210", "事件接入模块"],
                ["220", "事件归并模块"],
                ["230", "视觉复核模块"],
                ["240", "结果标准化模块"],
                ["250", "风险处理模块"],
                ["260", "派单路线模块"],
                ["270", "状态同步模块"],
                ["280", "审计归档模块"],
            ]},
            {"kind": "note", "text": "上传 CNIPA 时，图像本体使用附图目录中的 PNG，本文档作为图纸说明和编号索引保留在申请包中。"},
        ],
    )
    image_blocks = [
        {"kind": "image", "path": FIGURES / "图1-系统结构示意图.png", "caption": "图1  系统结构示意图", "width": 6.3},
        {"kind": "image", "path": FIGURES / "图2-事件处理流程示意图.png", "caption": "图2  事件处理流程示意图", "width": 6.3, "page_break_before": True},
        {"kind": "image", "path": FIGURES / "图3-证据帧复核流程示意图.png", "caption": "图3  证据帧复核流程示意图", "width": 6.3, "page_break_before": True},
        {"kind": "image", "path": FIGURES / "图4-事件状态审计示意图.png", "caption": "图4  事件状态审计示意图", "width": 6.3, "page_break_before": True},
    ]
    # Rebuild as a viewer-friendly docx with the actual figure PNGs.
    doc = Document()
    configure_doc(doc)
    add_title(doc, TITLE, "说明书附图")
    for block in image_blocks:
        if block.get("page_break_before"):
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(block["path"]), width=Inches(block["width"]))
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = c.add_run(block["caption"])
        set_run_font(run, size=10, bold=True)
    finish_doc(doc, OUT / "05_说明书附图.docx")


def build_abstract_figure_pack() -> None:
    blocks = [
        {"kind": "para", "text": "摘要附图采用图1作为代表图，用于说明本案的主闭环结构。", "indent": False},
        {"kind": "image", "path": OUT / "摘要附图.png", "caption": "摘要附图（对应图1）", "width": 6.3},
    ]
    write_markdown_case(OUT / "06_摘要附图.md", TITLE, "摘要附图说明", blocks)
    doc = Document()
    configure_doc(doc)
    add_title(doc, TITLE, "摘要附图")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(OUT / "摘要附图.png"), width=Inches(6.3))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = c.add_run("摘要附图（对应图1）")
    set_run_font(run, size=10, bold=True)
    finish_doc(doc, OUT / "06_摘要附图.docx")


def build_preliminary_search_notes_pack() -> None:
    blocks = [
        {"kind": "para", "text": "以下为初步查新笔记，目的是把本案与常见公开方向分开，供后续 CNIPA 数据库复核或代理人再检索。由于当前未接入专利数据库接口，以下结论均为预查新结论，不能替代正式检索。", "indent": False},
        {"kind": "heading", "text": "检索范围", "level": 2},
        {"kind": "bullets", "items": [
            "CNIPA / 国家知识产权局公开数据库",
            "Google Patents",
            "Google Scholar / 公开技术文章",
        ]},
        {"kind": "heading", "text": "检索关键词", "level": 2},
        {"kind": "table", "headers": ["关键词组", "用途"], "rows": [
            ["video alarm review dispatch patent", "寻找视频告警、复核和派单组合"],
            ["event dispatch video evidence patent", "寻找事件工单、证据帧和视频线索组合"],
            ["security patrol dispatch event review patent", "寻找巡防调度和复核闭环"],
            ["夜间商圈 安防 事件 归并", "寻找中文近似方案"],
        ]},
        {"kind": "heading", "text": "目前看到的近似方向", "level": 2},
        {"kind": "bullets", "items": [
            "单独的视频告警或图像识别系统：通常强调识别结果本身。",
            "单独的巡防派单或工单系统：通常强调调度，不一定有证据帧先保存机制。",
            "证据存储或复核系统：通常强调保存和查看，不一定有事件归并与就近派单闭环。",
            "安防消息推送系统：通常强调通知，不一定把复核结果回写到统一事件状态。",
        ]},
        {"kind": "heading", "text": "初步差异判断", "level": 2},
        {"kind": "para", "text": "本案最稳的差异点在于把“证据帧先落地、结构化结果异步上报、统一事件归并、视觉复核、就近派单、状态审计”组合成一条闭环链路，而不是把检测、派单、复核分成几个彼此割裂的模块。"},
        {"kind": "note", "text": "正式提交前仍需用 CNIPA 公开库和至少一轮人工专利检索再复核一次。"},
    ]
    write_markdown_case(OUT / "07_初步查新笔记.md", TITLE, "初步查新笔记", blocks)
    build_block_docx(OUT / "07_初步查新笔记.docx", TITLE, "初步查新笔记", blocks)


def build_request_confirmation_pack() -> None:
    blocks = [
        {"kind": "para", "text": "下表用于发明专利请求书信息确认。请把个人信息、权属信息和联系信息补齐后，再进入 CNIPA 网页填报。技术文件中不写身份证号、统一社会信用代码等敏感信息。", "indent": False},
        {"kind": "table", "headers": ["字段", "填写内容", "备注"], "rows": [
            ["发明名称", TITLE, "已预填"],
            ["申请类型", "发明专利", "已预填"],
            ["申请人名称", "待填写", "个人/企业/其他法人"],
            ["统一社会信用代码/证件号", "待填写", "只在请求书中填写"],
            ["申请人地址", "待填写", "含邮编"],
            ["联系人", "待填写", "建议为专利事务联系人"],
            ["联系电话", "待填写", ""],
            ["电子邮箱", "待填写", ""],
            ["发明人1", "待填写", "按实际贡献排序"],
            ["发明人2", "待填写", ""],
            ["其他发明人", "待填写", ""],
            ["是否要求优先权", "否 / 待确认", "如有在先申请需补充"],
            ["是否请求费用减缓", "待确认", "按实际资格决定"],
            ["是否提前公开", "按正常公开", "如需特别安排再改"],
            ["是否同时提出实质审查请求", "先不填 / 待确认", "可由代理人建议"],
        ]},
        {"kind": "heading", "text": "权属与合规提示", "level": 2},
        {"kind": "bullets", "items": [
            "若申请人为单位，需确认发明人归属和内部权利链条。",
            "若发明在公开演示、上线发布或销售中已出现，请先核对公开日期。",
            "如需保密，提交前应先和代理人确认是否存在涉密限制。",
        ]},
    ]
    write_markdown_case(OUT / "08_发明专利请求书信息确认表.md", TITLE, "发明专利请求书信息确认表", blocks)
    build_block_docx(OUT / "08_发明专利请求书信息确认表.docx", TITLE, "发明专利请求书信息确认表", blocks)


def build_difference_analysis_pack() -> None:
    blocks = [
        {"kind": "para", "text": "下面的差异化分析是给代理人和自己看的，不是正式法律结论。它的作用是把“看起来像常见安防系统”的部分和“更像本案核心发明点”的部分分开。", "indent": False},
        {"kind": "table", "headers": ["公开方向", "常见做法", "本案差异", "对权利要求的影响"], "rows": [
            ["视频告警系统", "直接输出告警或识别结果", "本案先保存证据帧，再异步发送结构化结果", "可把证据帧先落地写进独立特征"],
            ["巡防派单系统", "根据事件类型推送人员", "本案将位置、在线状态、职责和任务负载一起用于就近派单", "可把路线生成和人员筛选写进从属特征"],
            ["证据保存系统", "保存图片或录像供查看", "本案把证据和事件归并、复核、派单、状态审计连成闭环", "可把闭环链路写进独立权利要求"],
            ["通知推送系统", "仅负责提醒和消息流转", "本案把复核结果回写到风险等级和处置优先级", "可把风险回写写进权利要求"],
        ]},
        {"kind": "heading", "text": "最稳的保护点", "level": 2},
        {"kind": "bullets", "items": [
            "证据帧先保存，再异步上报结构化检测结果。",
            "多源事件统一入库并以 eventKey 做幂等归并。",
            "复核结果参与风险等级和派单优先级计算。",
            "处置状态与审计记录同步回写，形成闭环。",
        ]},
        {"kind": "heading", "text": "风险点", "level": 2},
        {"kind": "bullets", "items": [
            "如果把具体模型名写死，可能把范围压窄。",
            "如果未来查到更近的公开方案，需要调整独立权利要求的表述顺序。",
            "如果存在公开展示或已上线销售，需先确认法定公开影响。",
        ]},
    ]
    write_markdown_case(OUT / "09_查新与差异化分析.md", TITLE, "查新与差异化分析", blocks)
    build_block_docx(OUT / "09_查新与差异化分析.docx", TITLE, "查新与差异化分析", blocks)


def build_submission_checklist_pack() -> None:
    blocks = [
        {"kind": "para", "text": "以下为提交前清单。正式上传 CNIPA 时，通常还要在网页内逐项确认申请人、发明人和费用事项。", "indent": False},
        {"kind": "table", "headers": ["文件/事项", "状态", "备注"], "rows": [
            ["说明书", "已生成", "03_说明书.docx 或原始说明书.docx"],
            ["权利要求书", "已生成", "02_权利要求书.docx 或原始权利要求书.docx"],
            ["说明书摘要", "已生成", "04_说明书摘要.docx 或原始摘要.docx"],
            ["说明书附图", "已生成", "05_说明书附图.docx + 附图目录 PNG"],
            ["摘要附图", "已生成", "06_摘要附图.docx + 摘要附图.png"],
            ["技术交底书", "已生成", "01_技术交底书.docx"],
            ["初步查新笔记", "已生成", "07_初步查新笔记.docx"],
            ["请求书信息确认表", "已生成", "08_发明专利请求书信息确认表.docx"],
            ["查新与差异化分析", "已生成", "09_查新与差异化分析.docx"],
            ["提交文件清单", "已生成", "本文件"],
        ]},
        {"kind": "heading", "text": "上传顺序建议", "level": 2},
        {"kind": "bullets", "items": [
            "先核对请求书信息，再导入说明书、权利要求书和摘要。",
            "附图使用黑白 PNG，编号与说明书中的参考标号保持一致。",
            "正式提交前再复查申请人、发明人和公开日信息。",
        ]},
    ]
    write_markdown_case(OUT / "提交文件清单.md", TITLE, "提交文件清单", blocks)
    build_block_docx(OUT / "提交文件清单.docx", TITLE, "提交文件清单", blocks)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    REVIEW.mkdir(parents=True, exist_ok=True)
    build_description()
    build_claims()
    build_abstract()
    create_figures()
    build_core_doc_copies()
    build_overview_pack()
    build_technical_disclosure_pack()
    build_figures_companion_pack()
    build_abstract_figure_pack()
    build_preliminary_search_notes_pack()
    build_request_confirmation_pack()
    build_difference_analysis_pack()
    build_submission_checklist_pack()
    build_submission_template()
    build_evidence_manifest()
    print(f"generated: {OUT}")


if __name__ == "__main__":
    main()
