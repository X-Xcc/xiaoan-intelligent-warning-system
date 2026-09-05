from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:/xx/Desktop/查新委托单(1)(12)(1).docx")
OUTPUT = Path(r"D:/CICSIC/outputs/查新委托单_城脉智诊_已填写.docx")


PROJECT_CN = "城脉智诊-无监督驱动的地下管网分布式智能监测与边缘预警系统"
PROJECT_EN = (
    "CityPulse Intelligent Diagnosis: Unsupervised Distributed Monitoring "
    "and Edge Early Warning System for Underground Utility Networks"
)


TECHNICAL_POINTS = [
    (
        "本项目属于城市地下管网智能监测、物联网感知、多源传感数据融合、"
        "无监督异常检测、边缘计算和GIS运维预警技术领域。项目面向地下供水、"
        "排水、燃气、供热等管网运行中泄漏、微漏、爆管、堵塞、压力超限、"
        "低频小概率异常和多点并发异常难以及时发现的问题，解决传统人工巡检、"
        "单一传感器告警和中心端集中分析在样本标注依赖、实时性、定位精度、"
        "离线运行和处置闭环方面的不足。"
    ),
    (
        "系统通过分布式传感节点采集声学、压力、流量等多维运行数据，在边缘端完成"
        "数据清洗、时钟同步、特征提取、异常模式识别和本地预警；采用无监督学习与"
        "自学习适配机制，对不同管网工况下的泄漏、压力异常、堵塞、爆管、暗漏、"
        "微破损等异常进行识别，无需依赖大量预标注样本即可适配新增工况。"
    ),
    (
        "系统将异常识别结果与GIS管网图层、空间定位、历史异常数据、工单流程和多终端"
        "展示联动，形成异常感知、分级预警、空间定位、工单派发、处置核验、数据溯源"
        "和统计报表的闭环管理。边缘节点支持断网环境下本地独立运行、本地存储和联网后"
        "自动补传，适用于地下管网现场通信条件复杂的应用场景。"
    ),
    (
        "根据现有检验报告，本项目样品在异常事件感知精度、异常误报率、异常漏报率、"
        "多源传感数据融合准确率、端到端异常预警响应时间、异常点空间定位误差、"
        "工单自动派发准确率、处置结果核验准确率、边缘端离线运行和系统可靠性等方面"
        "已形成测试数据支撑。"
    ),
]


NOVELTY_POINTS = [
    (
        "1. 多源传感数据融合与无监督异常检测结合：对声学、压力、流量等地下管网"
        "运行数据进行融合分析，并通过无监督学习和自学习机制识别泄漏、微漏、堵塞、"
        "爆管、压力异常等多类工况。"
    ),
    (
        "2. 分布式边缘节点本地预警机制：在边缘端完成采集、清洗、时钟同步、异常识别、"
        "分级预警、本地存储和联网补传，支持断网条件下的连续监测和秒级响应。"
    ),
    (
        "3. 异常识别与GIS定位及工单闭环联动：将异常点定位结果与管网GIS图层、历史数据、"
        "工单自动派发、处置核验、多终端同步和数据溯源相结合，实现从监测到处置的闭环管理。"
    ),
]


SEARCH_TERMS = [
    "中文关键词：地下管网、地下管廊、城市管网、管线监测、压力管线、燃气管网、供水管网、排水管网。",
    "异常识别词：泄漏检测、微漏检测、爆管预警、堵塞识别、压力异常、流量异常、暗漏、小概率异常事件。",
    "算法与模型词：无监督学习、无监督异常检测、自学习、场景自适应、少样本、零样本、时序异常检测、数据驱动诊断。",
    "感知与融合词：多源传感、多传感器融合、声学传感、压力传感、流量监测、负压波、声波、光纤传感、数据融合。",
    "系统与运维词：边缘计算、分布式传感节点、边缘预警、离线运行、本地存储、断网补传、GIS管网、空间定位、工单派发、闭环处置、多终端同步。",
    "英文关键词：underground pipeline, utility network, leak detection, acoustic sensing, pressure monitoring, flow monitoring, multi-sensor fusion, unsupervised anomaly detection, self-learning, edge computing, GIS localization, work order dispatch, closed-loop maintenance.",
    "参考分类号：G01M 3/00、G01M 3/24、G01D 21/02、G08B 21/18、G06N 20/00、G06Q 10/20。分类号供检索扩展使用，正式检索时以查新机构分类复核为准。",
]


IP_STATUS = [
    "1. 现有资料显示，与本项目直接相关的产品检验材料为《城脉智诊-无监督驱动的地下管网分布式智能监测与边缘预警系统》检验报告，申请单位为沈阳大学，送检日期为2026年05月11日，检验日期为2026年05月11日至2026年05月18日。",
    "2. 现有项目资料未提供已申请或已授权专利号、软件著作权登记号、已发表论文清单或已获批立项文件编号；如委托方已有相关专利、软著、论文、项目立项或共有权属材料，建议在正式提交前补充名称、编号、申请人或作者、日期等信息。",
    "3. 如本项目存在对外公开展示、论文投稿、产品销售、试点部署或委托开发/合作开发情形，建议委托方补充公开日期、合作单位及权属说明，便于查新机构判断检索边界和成果归属。",
]


REFERENCES = [
    "1. CN115597790A，一种基于多源信息融合的供热管道泄漏检测与定位方法，公开日：2023-01-13，申请人：石家庄铁道大学。",
    "2. CN114627604B，一种地下管廊状态监测预警方法及系统，授权公告日：2023-08-15，申请人：相关公开专利文本。",
    "3. 钱志坚. 基于物联网与GIS的地下燃气管网监测系统研究. 测绘地理信息, 2019。",
    "4. Gong Y. et al. A Review on Gas Pipeline Leak Detection: Acoustic-Based Methods and Related Technologies. Information, 2025, 16(9):731。",
    "5. Zhang H. et al. An unsupervised leak detection method with aggregating strategies for long-term monitoring. Process Safety and Environmental Protection, 2023。",
    "6. Islam M. R. et al. Leak detection and localization in underground water supply systems. 2024。",
]


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def clear_cell(cell):
    for paragraph in list(cell.paragraphs):
        p = paragraph._element
        p.getparent().remove(p)


def set_cell_text(cell, text, size=10.5, bold=False):
    clear_cell(cell)
    paragraph = cell.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def append_paragraph(cell, text, size=10.0, bold=False):
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def append_section_content(cell, paragraphs):
    append_paragraph(cell, "填写内容：", size=10.5, bold=True)
    for item in paragraphs:
        append_paragraph(cell, item, size=10.0)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(SOURCE))
    table = doc.tables[0]

    set_cell_text(table.rows[0].cells[2], f"中文：{PROJECT_CN}", size=10.5)
    set_cell_text(table.rows[1].cells[2], f"英文：{PROJECT_EN}", size=10.0)
    set_cell_text(table.rows[2].cells[1], "沈阳大学", size=10.5)
    set_cell_text(table.rows[2].cells[2], "已填", size=10.5)
    set_cell_text(table.rows[3].cells[2], "沈阳市大东区望花南街21号", size=10.5)
    set_cell_text(table.rows[4].cells[2], "待确认", size=10.5)
    set_cell_text(table.rows[4].cells[4], "待确认", size=10.5)

    purpose = (
        "二、查新目的及范围 \n"
        "1.查新目的：(必填) \n"
        "○立项查新： ☐开题 ☐申报计划 ☐检查 ☐评估 ☐其他（请注明）\n"
        "○成果查新： ☐鉴定 ☐验收 ☐评估 ☐申报奖励 ☐其他（请注明）\n"
        "√产品查新 ☐申请专利 ☐标准查新 ☐其他（请注明）\n"
        "2.查新范围：(必填)\n"
        "√国内查新    ☐国内外查新"
    )
    set_cell_text(table.rows[6].cells[0], purpose, size=10.5)

    append_section_content(table.rows[7].cells[0], TECHNICAL_POINTS)
    append_section_content(table.rows[8].cells[0], NOVELTY_POINTS)
    append_section_content(table.rows[9].cells[0], SEARCH_TERMS)
    append_section_content(table.rows[10].cells[0], IP_STATUS)
    append_section_content(table.rows[11].cells[0], REFERENCES)

    for section in doc.sections:
        section.top_margin = section.top_margin

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
