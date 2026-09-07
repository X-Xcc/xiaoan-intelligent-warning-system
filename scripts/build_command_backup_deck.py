"""Fallback teaching pages, explicitly not browser acceptance screenshots."""
from pathlib import Path
import json
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "deliverables/command/fallback"
ASSETS = ROOT / "apps/dashboard/public/command"
DATA = json.loads((ROOT / "server/app/data/command/night_market_b1_b4_v1.json").read_text(encoding="utf-8"))
FONT = "C:/Windows/Fonts/msyh.ttc"


def write(draw, x, y, content, size=28, fill="#ffffff"):
    draw.text((x, y), content, font=ImageFont.truetype(FONT, size), fill=fill)


def paragraph(draw, text, x, y, width, size=30, fill="#e1edf6"):
    font = ImageFont.truetype(FONT, size)
    line = ""
    for character in text:
        if character == "\n":
            write(draw, x, y, line, size, fill)
            y += int(size * 1.7)
            line = ""
            continue
        if draw.textlength(line + character, font=font) > width and character not in "，。；：、！？":
            write(draw, x, y, line, size, fill)
            y += int(size * 1.7)
            line = character
        else:
            line += character
    if line:
        write(draw, x, y, line, size, fill)
    return y + int(size * 1.7)


def paste(image, filename, box):
    asset = Image.open(ASSETS / filename)
    fitted = ImageOps.contain(asset, (box[2], box[3]))
    image.paste(fitted, (box[0] + (box[2] - fitted.width) // 2, box[1] + (box[3] - fitted.height) // 2))


def page(stage, title):
    image = Image.new("RGB", (1920, 1080), "#0c447c")
    draw = ImageDraw.Draw(image)
    write(draw, 60, 38, f"{stage}   小安 · {title}", 48)
    write(draw, 1490, 55, "只读教学备用页", 26, "#f4d294")
    draw.line((60, 125, 1860, 125), fill="#5586b6", width=2)
    write(draw, 60, 151, DATA["title"] + " / " + DATA["intake"]["locationText"], 28, "#c4dcee")
    write(draw, 60, 991, f"场景V{DATA['version']} · 数据时间 {DATA['baseTime']}", 21, "#c4dcee")
    write(draw, 1360, 991, DATA["watermark"], 24, "#f0cc85")
    write(draw, 60, 1030, "教学备用图，不是浏览器验收截图；回放不代表本次业务已提交。", 18, "#bfd5e8")
    return image, draw


def build():
    DEST.mkdir(parents=True, exist_ok=True)
    images = []
    image, d = page("B1", "接警工单")
    write(d, 70, 250, DATA["intake"]["speakerName"] + " · 报警文本", 34)
    paragraph(d, DATA["intake"]["transcript"], 70, 320, 1050, 38)
    write(d, 1250, 250, "消费纠纷升级风险", 36, "#f4c16e")
    paragraph(d, "伤情、危险物品与现场冲突情况待人工核实。", 1250, 320, 575, 30)
    write(d, 70, 610, "关联警情 · 2起", 36)
    for i, item in enumerate(DATA["relatedAlerts"]):
        x = 70 + i * 900
        d.rectangle((x, 680, x + 840, 890), fill="#185fa5")
        write(d, x + 25, 705, item["eventId"], 30, "#f2d595")
        write(d, x + 25, 760, item["locationText"], 29)
        paragraph(d, "；".join(item["basis"]), x + 25, 813, 790, 26)
    images.append(("b1", image))
    image, d = page("B2", "导航派警")
    paste(image, "night-market-map.png", (60, 245, 1250, 680))
    write(d, 1360, 265, "135快反组", 44)
    write(d, 1360, 365, "至入口约3分钟", 41, "#f4d294")
    write(d, 1360, 443, "入口至现场步行200米", 33)
    paragraph(d, DATA["route"]["notice"], 1360, 550, 500, 30)
    paragraph(d, "确认派警建议与下达派警是两个独立人工动作。", 1360, 770, 500, 27)
    images.append(("b2", image))
    image, d = page("B3", "盘查核验")
    write(d, 185, 330, DATA["person"]["name"], 84)
    write(d, 185, 480, DATA["person"]["archiveNo"], 36, "#f4d294")
    write(d, 185, 560, "教学档案查询", 31)
    y = 300
    for basis in DATA["person"]["basis"]:
        y = paragraph(d, basis, 900, y, 850, 36) + 35
    paragraph(d, "规则匹配分：不提供\n模型置信度：不提供", 900, 560, 850, 30)
    paragraph(d, "核验结果为线索，需人工确认；不生成身份或案件结论。", 185, 785, 1540, 39, "#f4d294")
    images.append(("b3", image))
    image, d = page("B4", "物证同步")
    for index, material in enumerate(DATA["materials"]):
        write(d, 65, 260 + index * 145, f"{index + 1:02}  {material['name']}", 35, "#f4d294")
        paragraph(d, material["description"], 115, 317 + index * 145, 480, 24)
    paste(image, "receipts.png", (650, 230, 1200, 400))
    paste(image, "notebook.png", (650, 640, 570, 310))
    paste(image, "phone.png", (1250, 640, 580, 310))
    images.append(("b4", image))
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    for stage, image in images:
        destination = DEST / f"{stage}-teaching.png"
        image.save(destination, optimize=True)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(destination), 0, 0, width=deck.slide_width, height=deck.slide_height)
        slide.notes_slide.notes_text_frame.text = f"场景V1.0。来源：night_market_b1_b4_v1.json。教学备用图，非浏览器验收截图。"
    deck.save(DEST / "小安接处警-教学备用页.pptx")
    contact = Image.new("RGB", (1920, 1080), "#fff")
    for i, (_, image) in enumerate(images):
        contact.paste(image.resize((960, 540)), ((i % 2) * 960, (i // 2) * 540))
    contact.save(DEST / "contact-sheet.png")
    print("Generated four teaching PNG pages, one contact sheet and the fallback PPTX.")


if __name__ == "__main__":
    build()
