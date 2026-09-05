from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\xx\Desktop\查新委托单(1)(12)(1).docx")
OUTPUT = Path(r"D:\CICSIC\_artifact_work\查新委托单-烟火哨兵已填写-定稿.docx")
AUDIT = Path(r"D:\CICSIC\_artifact_work\查新委托单-烟火哨兵填写核对.json")


PROJECT_CN = "烟火哨兵-夜市治安风险智能预警与联控平台"
PROJECT_EN = (
    "Yanhuo Sentinel - Smart Early Warning and Coordinated Response Platform "
    "for Night Market Public Security Risks"
)
UNIT = "江西司法警官职业学院"
ADDRESS = "江西省南昌市昌北经济技术开发区长青路76号"
CONTACT = "徐俊"
PHONE = "13456739866"


TECHNICAL_POINTS = (
    "本项目属于夜间商圈治安风险智能预警、视频分析、事件数据处理、移动通信、巡防调度和安全指挥技术领域，"
    "面向夜市场景中群众求助、商户上报、视频检测、巡防反馈及人工录入等多源安全事件并发输入时，容易出现信息分散、"
    "重复建单、证据与事件脱离、复核结果难以参与调度以及处置过程难以追溯的问题，构建群众端、视频检测端、巡防端、"
    "指挥端和后端事件库协同的风险预警与联控平台。\n"
    "系统对群众、商户、视频、巡防和人工录入等来源的数据进行统一结构化解析，形成包含来源、位置、时间、事件类型、"
    "风险等级、证据索引和处置状态的事件记录，并依据 eventKey 或同类唯一标识进行事件归并和幂等处理，避免同一风险重复建单。\n"
    "视频检测侧对采样帧执行识别，在满足预设事件条件时先保存证据帧，再异步上报包含动作标签、人员数量、相机标识、"
    "时间戳、帧数、帧率和证据帧路径的结构化结果；画面复核模块结合证据帧与检测上下文输出风险等级、人数估计、场景摘要"
    "和处置建议，复核服务异常时仍保留本地检测结果供人工确认。\n"
    "派单模块综合事件位置、巡防人员当前位置、在线状态、职责范围和任务负载，生成就近处置人员与响应路线；平台将提交、"
    "派单、接收、到达、处理中、完成等状态同步至指挥端和巡防端，并在审计记录中保存风险变化、派单结果、处置结果和状态变化，"
    "形成从事件发现、证据固定、风险复核、协同处置到归档复盘的闭环。"
)

NOVELTY_POINTS = (
    "1. 多源安全事件统一结构化入库、事件归并与幂等处理：将群众、商户、视频、巡防和人工录入等至少两类来源归一为统一事件结构，"
    "并依据事件标识进行归并，减少同一风险的重复建单。\n"
    "2. 证据帧先保存、检测结果异步上报与结构化视觉复核结合的风险分级机制：在视频检测满足事件条件时先固定证据帧，再上报结构化结果，"
    "并基于证据帧和检测上下文输出风险等级、场景摘要与处置建议。\n"
    "3. 基于位置、在线状态、职责和任务负载的就近派单及全流程审计机制：生成推荐处置人员和响应路线，同步记录提交、派单、接收、到达、"
    "处理、完成状态及相应审计信息。"
)

SEARCH_TERMS = (
    "中文检索词：夜市治安；夜间商圈安全治理；治安风险预警；多源安全事件；事件统一入库；事件归并；幂等处理；视频检测；"
    "证据帧；异步上报；结构化视觉复核；风险分级；巡防调度；就近派单；路线推荐；状态同步；审计日志；闭环处置。\n"
    "英文检索词：night market public security; night-time commercial district safety; public security risk early warning;"
    " multi-source safety events; unified event ingestion; event deduplication; idempotent event processing; video detection;"
    " evidence frame; asynchronous reporting; structured visual review; risk grading; patrol dispatch; nearest-unit dispatch;"
    " route recommendation; status synchronization; audit trail; closed-loop response.\n"
    "英文项目名称：" + PROJECT_EN
)

IP_STATUS = (
    "现有项目资料包含技术交底书、专利申请材料草案、系统实现文档和交付说明。截止本委托单填写时，未提供与本项目直接相关的已公开专利号、"
    "软件著作权登记号、已发表论文或已获批立项编号；如有新增知识产权、论文、立项或共有权属材料，以委托方补充清单为准。"
)

REFERENCES = (
    "1. 《烟火哨兵项目实现文档》，项目技术资料，2026年。\n"
    "2. 《烟火哨兵交付说明》，项目技术资料，2026年。\n"
    "3. 《一种面向夜间商圈的多源安全事件统一入库、证据帧复核与就近派单闭环方法及系统》技术交底书，项目技术资料，2026年。\n"
    "4. GB/T 28181-2022，《公共安全视频监控联网系统信息传输、交换、控制技术要求》。\n"
    "5. GB/T 35273-2020，《信息安全技术 个人信息安全规范》。"
)


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def clear_cell(cell) -> None:
    for paragraph in list(cell.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)


def add_paragraph(cell, text: str, *, size: float = 10.0, bold: bool = False,
                  align: WD_ALIGN_PARAGRAPH | None = None) -> None:
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def set_cell_text(cell, text: str, *, size: float = 10.0,
                  bold: bool = False, align: WD_ALIGN_PARAGRAPH | None = None) -> None:
    clear_cell(cell)
    for line in str(text).split("\n"):
        add_paragraph(cell, line, size=size, bold=bold, align=align)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_section_cell(cell, title: str, body: str, *, body_size: float = 8.8) -> None:
    clear_cell(cell)
    add_paragraph(cell, title, size=10.5, bold=True)
    for line in body.split("\n"):
        add_paragraph(cell, line, size=body_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def normalize_table_rows(table) -> None:
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_form() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)

    doc = Document(str(OUTPUT))
    if len(doc.tables) != 1:
        raise ValueError(f"模板表格数量异常: {len(doc.tables)}")
    table = doc.tables[0]
    if len(table.rows) != 13 or len(table.columns) != 5:
        raise ValueError(f"模板表格结构异常: {len(table.rows)} 行 x {len(table.columns)} 列")
    normalize_table_rows(table)

    set_cell_text(table.rows[0].cells[2], f"中文：{PROJECT_CN}", size=10.6)
    set_cell_text(table.rows[1].cells[2], f"英文：{PROJECT_EN}", size=8.65)
    set_cell_text(table.rows[2].cells[2], UNIT, size=10.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.rows[3].cells[2], ADDRESS, size=10.2)
    set_cell_text(table.rows[4].cells[2], CONTACT, size=10.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.rows[4].cells[4], PHONE, size=10.8, align=WD_ALIGN_PARAGRAPH.CENTER)

    set_cell_text(
        table.rows[6].cells[0],
        "二、查新目的及范围\n"
        "1. 查新目的：\n"
        "○立项查新：□开题  □申报计划  □检查  □评估  □其他（请注明）\n"
        "○成果查新：□鉴定  □验收  □评估  □申报奖励  □其他（请注明）\n"
        "√产品查新  □申请专利  □标准查新  □其他（请注明）\n"
        "2. 查新范围：□国内查新    √国内外查新",
        size=8.6,
    )
    set_section_cell(table.rows[7].cells[0], "三、查新项目的科学技术要点", TECHNICAL_POINTS, body_size=8.45)
    set_section_cell(table.rows[8].cells[0], "四、查新点", NOVELTY_POINTS, body_size=8.55)
    set_section_cell(table.rows[9].cells[0], "五、参考检索词及其解释", SEARCH_TERMS, body_size=8.35)
    set_section_cell(table.rows[10].cells[0], "六、知识产权及已发表论文情况", IP_STATUS, body_size=8.8)
    set_section_cell(table.rows[11].cells[0], "七、参考文献", REFERENCES, body_size=8.55)

    core = doc.core_properties
    core.title = "科技查新委托单"
    core.subject = PROJECT_CN
    core.keywords = "科技查新;烟火哨兵;夜市治安;风险预警;联控"
    doc.save(str(OUTPUT))


def read_back_and_audit() -> dict:
    doc = Document(str(OUTPUT))
    table = doc.tables[0]
    values = {
        "中文项目名称": table.rows[0].cells[2].text,
        "英文项目名称": table.rows[1].cells[2].text,
        "委托单位": table.rows[2].cells[2].text,
        "收件地址": table.rows[3].cells[2].text,
        "收件人": table.rows[4].cells[2].text,
        "手机": table.rows[4].cells[4].text,
        "查新目的及范围": table.rows[6].cells[0].text,
        "科学技术要点": table.rows[7].cells[0].text,
        "查新点": table.rows[8].cells[0].text,
        "参考检索词": table.rows[9].cells[0].text,
        "知识产权及论文": table.rows[10].cells[0].text,
        "参考文献": table.rows[11].cells[0].text,
    }
    all_text = "\n".join(values.values())
    forbidden_placeholders = ["必填", "待确认", "请填写", "国内查新必填", "国内外查新必填"]
    forbidden_residue = ["城脉智诊", "沈阳大学", "地下管网"]
    unresolved = {
        "required_placeholders": [item for item in forbidden_placeholders if item in all_text],
        "prior_project_residue": [item for item in forbidden_residue if item in all_text],
    }
    purpose = values["查新目的及范围"]
    selected = {
        "产品查新": "√产品查新" in purpose,
        "国内外查新": "√国内外查新" in purpose,
    }
    audit = {
        "source_template": str(SOURCE),
        "output": str(OUTPUT),
        "filled_fields": values,
        "selected_options": selected,
        "unresolved": unresolved,
        "checks": {
            "template_structure": len(table.rows) == 13 and len(table.columns) == 5,
            "only_requested_options_marked": selected == {"产品查新": True, "国内外查新": True},
            "no_required_placeholders": not unresolved["required_placeholders"],
            "no_prior_project_residue": not unresolved["prior_project_residue"],
            "contact_phone_digits": PHONE == values["手机"],
        },
        "address_verification": {
            "normalized_address": ADDRESS,
            "basis": "按委托信息填写，提交前由委托方确认",
        },
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    return audit


if __name__ == "__main__":
    build_form()
    result = read_back_and_audit()
    print(json.dumps({
        "output": str(OUTPUT),
        "audit": str(AUDIT),
        "checks": result["checks"],
    }, ensure_ascii=False))
